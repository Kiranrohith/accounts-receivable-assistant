import sys
from docx import Document
from PIL import Image, ImageDraw, ImageFont
out = sys.argv[1]
d = Document()
d.add_heading("TAX INVOICE", 1)
d.add_paragraph("Invoice No: INV-2026-0042   Date: 2026-10-01   Due: 2026-10-31")
d.add_paragraph("Customer: Acme Traders Pvt Ltd   GSTIN: 29ABCDE1234F1Z5")
t = d.add_table(rows=2, cols=3); t.style = "Table Grid"
for i, h in enumerate(("Description", "Qty", "Amount")): t.rows[0].cells[i].text = h
for i, v in enumerate(("Widget A", "10", "50,000.00")): t.rows[1].cells[i].text = v
d.add_paragraph("Total INR 50,000.00")
d.save(f"{out}/inv.docx")
img = Image.new("RGB", (900, 400), "white"); dr = ImageDraw.Draw(img)
try: f = ImageFont.truetype("arial.ttf", 28)
except Exception: f = ImageFont.load_default()
for n, line in enumerate(("TAX INVOICE", "Invoice No: INV-2026-0043", "Customer: Acme Traders Pvt Ltd", "Total INR 75,000.00")):
    dr.text((40, 40 + n * 60), line, fill="black", font=f)
img.save(f"{out}/inv.png")
open(f"{out}/inv.xyz", "w").write("x")
