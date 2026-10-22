import asyncio
import re
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, "src")
import httpx
from docx import Document
from sqlalchemy import text

import core.services.documents.classification as cls_mod
import core.services.documents.intake as intake
import core.services.invoices.extraction as inv_ext
import core.services.purchase_orders.extraction as po_ext
from api.rest.app import app
from config.settings import get_settings
from core.services.documents.classification import ClassificationOutcome, ExtractionError
from data.clients.db import SessionLocal, engine
from schemas.documents.classification import DocClassification, DocumentClass

s = get_settings()
assert "ar_in" in s.database_url, "refusing to run against a non-test database"
OURS = "33AAHCG2210A1Z7"
results = []


def check(name, cond, extra=""):
    results.append(bool(cond))
    print(f"RESULT {'PASS' if cond else 'FAIL'} {name} {extra}"[:340])


# ---- count every LLM call by schema, and force classifier verdicts on marked documents
calls: list[str] = []


def wrap(mod):
    real = mod.run_structured_extraction

    async def counted(schema, messages):
        calls.append(schema.__name__)
        return await real(schema, messages)

    mod.run_structured_extraction = counted


for m in (cls_mod, inv_ext, po_ext):
    wrap(m)

real_classify = intake.classify_document


async def classify(markdown):
    m = re.search(r"FORCE:(\w+):?(\d+)?", markdown)
    if m:
        if m.group(1) == "boom":
            raise ExtractionError("model unavailable")
        verdict = DocClassification(
            doc_class=DocumentClass(m.group(1)), confidence=int(m.group(2)), reason="forced for the test"
        )
        return ClassificationOutcome(verdict, "fake-model", 1)
    return await real_classify(markdown)


intake.classify_document = classify


# ---- documents
def para(d, text_):
    d.add_paragraph(text_)


def invoice(path, *, number="INV-IN-001", po=None, lines=(("Laptop", "2", "25000.00"),), total="50000.00",
            marker=None, title="TAX INVOICE"):
    d = Document()
    d.add_heading(title, 1)
    para(d, f"Genworx Technologies Pvt Ltd\n12 Anna Salai, Chennai 600002\nGSTIN: {OURS}")
    para(d, (f"Invoice No: {number}\n" if number else "") + "Invoice Date: 05/10/2026\nDue Date: 04/11/2026")
    if po:
        para(d, f"Customer PO No: {po}")
    para(d, "Bill To:\nZenith Softlabs Pvt Ltd\nBengaluru\nGSTIN: 33AABCZ1234F1Z5")
    t = d.add_table(rows=1 + len(lines), cols=3)
    t.style = "Table Grid"
    for i, h in enumerate(("Description", "Qty", "Amount (INR)")):
        t.rows[0].cells[i].text = h
    for r, row in enumerate(lines, start=1):
        for i, v in enumerate(row):
            t.rows[r].cells[i].text = v
    para(d, f"Total INR {total}")
    if marker:
        para(d, marker)
    d.save(path)


def purchase_order(path, *, number="PO-IN-1", lines=(("Laptop", "10", "25000.00"),), total="250000.00", marker=None):
    d = Document()
    d.add_heading("PURCHASE ORDER", 1)
    para(d, "Zenith Softlabs Pvt Ltd\nBengaluru\nGSTIN: 33AABCZ1234F1Z5")
    para(d, (f"PO No: {number}\n" if number else "") + "PO Date: 02/10/2026")
    para(d, f"Supplier: Genworx Technologies Pvt Ltd, GSTIN {OURS}")
    t = d.add_table(rows=1 + len(lines), cols=3)
    t.style = "Table Grid"
    for i, h in enumerate(("Description", "Qty", "Unit price (INR)")):
        t.rows[0].cells[i].text = h
    for r, row in enumerate(lines, start=1):
        for i, v in enumerate(row):
            t.rows[r].cells[i].text = v
    para(d, f"Order total INR {total}")
    if marker:
        para(d, marker)
    d.save(path)


def contract(path):
    d = Document()
    d.add_heading("MASTER SERVICES AGREEMENT", 1)
    para(d, "This Agreement is made on 1 September 2026 between Genworx Technologies Pvt Ltd (the Supplier) "
            "and Zenith Softlabs Pvt Ltd (the Client).")
    para(d, "1. Term. This Agreement continues for 24 months unless terminated earlier by either party with 60 days' notice.")
    para(d, "2. Confidentiality. Each party shall keep the other's confidential information secret.")
    para(d, "3. Governing law. This Agreement is governed by the laws of India. Disputes go to the courts at Chennai.")
    para(d, "Signed for the Supplier: ____________    Signed for the Client: ____________")
    d.save(path)


def receipt(path):
    d = Document()
    d.add_heading("PAYMENT RECEIPT", 1)
    para(d, "Received with thanks from Zenith Softlabs Pvt Ltd the sum of INR 50,000.00 by NEFT, reference UTR12345678, "
            "towards invoice INV-0099 dated 01/09/2026.\nReceipt No: RCT-0042   Date: 10/10/2026")
    para(d, "Genworx Technologies Pvt Ltd")
    d.save(path)


def quotation(path):
    d = Document()
    d.add_heading("QUOTATION", 1)
    para(d, f"Genworx Technologies Pvt Ltd, GSTIN {OURS}\nQuote No: Q-2026-118   Valid until 30/11/2026\nTo: Zenith Softlabs Pvt Ltd")
    t = d.add_table(rows=2, cols=3)
    t.style = "Table Grid"
    for i, h in enumerate(("Description", "Qty", "Estimated price (INR)")):
        t.rows[0].cells[i].text = h
    for i, v in enumerate(("Annual support", "1", "120000.00")):
        t.rows[1].cells[i].text = v
    para(d, "This is an estimate, not a bill. Prices are subject to change.")
    d.save(path)


async def main():
    import testsupport
    worker = await testsupport.start_worker()
    tmp = Path(tempfile.mkdtemp())
    invoice(tmp / "inv_ok.docx", po="PO-IN-1")
    purchase_order(tmp / "po_ok.docx")
    invoice(tmp / "inv_nonumber.docx", number=None)
    purchase_order(tmp / "po_nonumber.docx", number=None)
    contract(tmp / "contract.docx")
    receipt(tmp / "receipt.docx")
    quotation(tmp / "quotation.docx")
    purchase_order(tmp / "f_po_low.docx", number="PO-IN-2", marker="FORCE:po:60")
    invoice(tmp / "f_inv_unclear.docx", number="INV-IN-002", marker="FORCE:unclear:90")
    invoice(tmp / "f_inv_low_reject.docx", number="INV-IN-003", marker="FORCE:invoice:50")
    purchase_order(tmp / "f_po_as_inv.docx", number="PO-IN-3", marker="FORCE:unclear:80")
    purchase_order(tmp / "f_nonum_po.docx", number=None, marker="FORCE:invoice:30")
    invoice(tmp / "f_race.docx", number="INV-IN-004", marker="FORCE:unclear:70")
    invoice(tmp / "f_boom.docx", number="INV-IN-005", marker="FORCE:boom")
    invoice(tmp / "f_confident_other.docx", number="INV-IN-006", marker="FORCE:other:90")
    invoice(tmp / "f_fa.docx", number="INV-IN-007", marker="FORCE:unclear:60")
    invoice(tmp / "f_choose_twice.docx", number="INV-IN-008", marker="FORCE:invoice:74")
    invoice(tmp / "f_edge75.docx", number="INV-IN-009", marker="FORCE:invoice:75")

    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t", timeout=900) as c:
        testsupport.install(c)
        login = lambda e, p: c.post("/auth/login", json={"email": e, "password": p})
        ADM = {"Authorization": f"Bearer {(await login(s.seed_admin_email, s.seed_admin_password)).json()['access_token']}"}
        await c.post("/admin/users", headers=ADM, json={"name": "Fa Two", "email": "fa.two@example.com", "password": "pass12345", "role": "FA"})
        FA = {"Authorization": f"Bearer {(await login('fa.two@example.com', 'pass12345')).json()['access_token']}"}
        cust = lambda **kw: {"contact": "Contact", "country": "India", "billing_currency": "INR", **kw}
        assert (await c.post("/customers", headers=ADM, json=cust(customer_name="Zenith Softlabs Pvt Ltd", email="ap@zenith.com", gstin="33AABCZ1234F1Z5", deducts_tds=True))).status_code == 201

        async def q(sql, **p):
            async with SessionLocal() as db:
                return (await db.execute(text(sql), p)).all()

        async def jobs(sid):
            return await q("select job_type, state from jobs where entity_id=:i order by created_at", i=sid)

        async def wait(sid, n=1):
            for _ in range(600):
                js = await jobs(sid)
                if len(js) >= n and js[-1][1] in ("Succeeded", "Failed"):
                    return js
                await asyncio.sleep(0.5)
            return await jobs(sid)

        async def status(sid, H=ADM):
            r = await c.get(f"/documents/uploads/{sid}", headers=H)
            assert r.status_code == 200, r.text
            return r.json()

        async def steps(sid):
            return [r[0] for r in await q("select step_name from step_log where entity_id=:i order by occurred_at, id", i=sid)]

        async def up(name, H=ADM):
            before = len(calls)
            r = await c.post("/documents/upload", headers=H, files={"file": (name, (tmp / name).read_bytes())})
            assert r.status_code == 202, (name, r.text)
            sid = r.json()["source_document_id"]
            await wait(sid)
            return sid, calls[before:]

        async def file_exists(sid):
            row = (await q("select file_path from source_documents where id=:i", i=sid))[0][0]
            return bool(row) and Path(row).exists()

        # ===================================================== real LLM: invoice
        sid, used = await up("inv_ok.docx")
        st = await status(sid)
        check("invoice: created, status label", st["status"] == "Invoice created" and st["disposition"] == "invoice_created", str(st["status"]))
        check("invoice: classifier then invoice extraction only", used == ["DocClassification", "ExtractedInvoice"], str(used))
        check("invoice: detected type + confidence on the upload", st["detected_type"] == "invoice" and (st["detected_confidence"] or 0) >= 75, f"{st['detected_type']} {st['detected_confidence']}")
        check("invoice: steps parse, classify, extract, validate, route", (await steps(sid))[:3] == ["parse", "classify", "extract"] and "route" in await steps(sid), str(await steps(sid)))
        cl = next(x for x in st["steps"] if x["step_name"] == "classify")
        check("classify step records class, confidence, threshold 75 and route", cl["output_summary"]["doc_class"] == "invoice" and cl["output_summary"]["threshold"] == 75 and cl["output_summary"]["route"] == "invoice" and cl["model_used"], str(cl["output_summary"]))
        raw = (await q("select raw_extraction_json from source_documents where id=:i", i=sid))[0][0]
        check("raw_extraction_json keeps the classification beside the fields", raw["classification"]["doc_class"] == "invoice" and raw["invoice_number"] == "INV-IN-001")
        check("invoice quoting a PO number is NOT taken for a PO", (await q("select doc_kind from source_documents where id=:i", i=sid))[0][0] == "invoice" and (await q("select count(*) from purchase_orders"))[0][0] == 0)
        inv_sid = sid

        # ===================================================== real LLM: purchase order
        sid, used = await up("po_ok.docx")
        js = await jobs(sid)
        st = await status(sid)
        check("PO: created through the invoice upload entry", st["status"] == "Purchase order created" and st["disposition"] == "po_created", str(st["status"]))
        check("PO: classifier then PO extraction only", used == ["DocClassification", "ExtractedPurchaseOrder"], str(used))
        check("PO: document switched to kind po (the job is still the one intake job)", (await q("select doc_kind from source_documents where id=:i", i=sid))[0][0] == "po" and js[-1][0] == "document_intake", str(js))
        check("PO: retro-linked the waiting invoice INV-IN-001, which passes", [(x["invoice_number"], x["passed"]) for x in st["linked_invoices"]] == [("INV-IN-001", True)], str(st["linked_invoices"]))
        listing = (await c.get("/documents/uploads", headers=ADM)).json()
        row = next((u for u in listing if u["source_document_id"] == sid), None)
        check("PO: listed in the one upload list as document_kind po", row is not None and row["document_kind"] == "po" and row["status"] == "Purchase order created", str(row and row["document_kind"]))
        check("PO: step log has classify with doc_class po", any(x["step_name"] == "classify" and x["output_summary"]["doc_class"] == "po" for x in st["steps"]))

        check("document_kind: invoice for the invoice, unknown for waiting/rejected ones (checked below)", st["document_kind"] == "po" and (await status(inv_sid))["document_kind"] == "invoice", st["document_kind"])
        check("old PO upload/status URLs are gone", (await c.post("/purchase-orders/upload", headers=ADM, files={"file": ("x.docx", b"x")})).status_code in (404, 405) and (await c.get("/purchase-orders/uploads", headers=ADM)).status_code in (404, 405, 422))
        check("old invoice upload/status URLs are gone", (await c.post("/invoices/upload", headers=ADM, files={"file": ("x.docx", b"x")})).status_code in (404, 405) and (await c.get("/invoices/uploads", headers=ADM)).status_code in (404, 405, 422))

        # ===================================================== real LLM: confidently neither
        others = {}
        for name in ("contract.docx", "receipt.docx", "quotation.docx"):
            sid, used = await up(name)
            others[name] = sid
            st = await status(sid)
            check(f"{name}: rejected, not extracted (only the classifier ran; a retried attempt may repeat it)", st["status"] == "Rejected" and set(used) == {"DocClassification"}, f"{st['status']} {used} {st.get('detected_type')} {st.get('detected_confidence')}")
            check(f"{name}: reason names 'neither an invoice nor a purchase order'", any("neither an invoice nor a purchase order" in r for r in st["disposition_reasons"]), str(st["disposition_reasons"]))
            check(f"{name}: file kept on disk, no extract step logged", await file_exists(sid) and "extract" not in await steps(sid))
            check(f"{name}: job Succeeded (a business outcome, not a failure)", (await jobs(sid))[-1][1] == "Succeeded")
            check(f"{name}: document_kind unknown, classifier verdict shown", st["document_kind"] == "unknown" and st["detected_type"] == "other" and (await c.get(f"/documents/uploads/{sid}", headers=ADM)).json()["classification"]["doc_class"] == "other")
        rej = (await c.get("/invoices/rejected", params={"kind": "other"}, headers=ADM)).json()
        check("rejected list kind=other holds the three, kind 'other'", {r["source_document_id"] for r in rej} == set(others.values()) and all(r["kind"] == "other" and r["file_kept"] for r in rej), str([r["kind"] for r in rej]))
        rej_inv = (await c.get("/invoices/rejected", headers=ADM)).json()
        check("default (invoice) rejected list excludes them", not ({r["source_document_id"] for r in rej_inv} & set(others.values())))
        r = await c.post("/documents/upload", headers=ADM, files={"file": ("contract.docx", (tmp / "contract.docx").read_bytes())})
        check("same file again -> 409 duplicate (hash kept, file kept)", r.status_code == 409, r.text[:120])

        # ===================================================== second net: validation
        sid, used = await up("inv_nonumber.docx")
        st = await status(sid)
        check("invoice with no invoice number: ends in rejected list via validation", st["status"] == "Rejected" and any("Invoice number is missing" in r for r in st["disposition_reasons"]), f"{st['status']} {st['disposition_reasons']}")
        check("...not in the invoices table, file kept (nothing deleted)", (await q("select count(*) from invoices where source_document_id=:i", i=sid))[0][0] == 0 and await file_exists(sid))
        sid, used = await up("po_nonumber.docx")
        st = await status(sid)
        check("PO with no PO number: rejected by PO validation", st["status"] == "Rejected" and any("PO number" in r for r in st["disposition_reasons"]), f"{st['status']} {st['disposition_reasons']}")

        # ===================================================== doc-type review queue
        calls_before = len(calls)
        sid, used = await up("f_po_low.docx")
        st = await status(sid)
        check("low-confidence PO (60) -> Needs review, review_kind document_type", st["status"] == "Needs review" and st["review_kind"] == "document_type" and st["detected_type"] == "po" and st["detected_confidence"] == 60, f"{st['status']} {st['review_kind']}")
        check("...no field extraction ran (zero LLM calls)", used == [], str(used))
        rid = st["review_id"]
        check("waiting for a type decision: document_kind unknown, no extraction shown", st["document_kind"] == "unknown" and (await c.get(f"/documents/uploads/{sid}", headers=ADM)).json()["extraction"] is None)
        queue = (await c.get("/documents/type-reviews", headers=ADM)).json()
        item = next((x for x in queue if x["review_id"] == rid), None)
        check("review queue lists it with suggestion, confidence and reasons", item and item["suggested_type"] == "po" and item["confidence"] == 60 and any("60%" in r for r in item["reasons"]), str(item))
        check("not visible in the invoice review list, and 404 there", rid not in [x["review_id"] for x in (await c.get("/invoices/reviews", headers=ADM)).json()] and (await c.get(f"/invoices/reviews/{rid}", headers=ADM)).status_code == 404)
        check("not visible in the PO review list", rid not in [x["review_id"] for x in (await c.get("/purchase-orders/reviews", headers=ADM)).json()])
        r = await c.post(f"/documents/type-reviews/{rid}/choose", headers=ADM, json={"doc_type": "po"})
        check("choose PO -> 200 processing", r.status_code == 200 and r.json()["processing"] is True and r.json()["chosen_type"] == "po", r.text[:150])
        await wait(sid, 2)
        st = await status(sid)
        check("picked PO: PO validation + creation ran -> Purchase order created", st["status"] == "Purchase order created", f"{st['status']} {st['disposition_reasons']}")
        human = [x for x in st["steps"] if x["step_name"] == "classify" and x["output_summary"].get("source") == "human"]
        check("step log records the human decision (class po, decided_by)", len(human) == 1 and human[0]["output_summary"]["doc_class"] == "po" and human[0]["output_summary"]["decided_by"], str(human))
        check("review item Approved, queue now empty of it", rid not in [x["review_id"] for x in (await c.get("/documents/type-reviews", headers=ADM)).json()] and (await c.get(f"/documents/type-reviews/{rid}", headers=ADM)).json()["status"] == "Approved")
        check("extraction ran once, after the choice", calls[calls_before:] == ["ExtractedPurchaseOrder"], str(calls[calls_before:]))
        r = await c.post(f"/documents/type-reviews/{rid}/choose", headers=ADM, json={"doc_type": "invoice"})
        check("deciding again -> 409", r.status_code == 409, r.text[:100])

        sid, used = await up("f_inv_unclear.docx")
        st = await status(sid)
        check("'unclear' even at 90 -> review (cannot be separated)", st["status"] == "Needs review" and st["review_kind"] == "document_type" and st["detected_type"] == "unclear", str(st["status"]))
        r = await c.post(f"/documents/type-reviews/{st['review_id']}/choose", headers=ADM, json={"doc_type": "invoice", "reason": "It is an invoice"})
        await wait(sid, 2)
        st = await status(sid)
        check("picked invoice -> invoice created with the usual flow", st["status"] == "Invoice created" and st["invoice_number"] == "INV-IN-002", f"{st['status']} {st['disposition_reasons']}")
        check("decision reason stored", (await c.get(f"/documents/type-reviews/{r.json()['review_id']}", headers=ADM)).json()["decision_reason"] == "It is an invoice")

        sid, used = await up("f_inv_low_reject.docx")
        st = await status(sid)
        rid = st["review_id"]
        r = await c.post(f"/documents/type-reviews/{rid}/reject", headers=ADM, json={})
        check("reject needs a reason (422)", r.status_code == 422)
        r = await c.post(f"/documents/type-reviews/{rid}/reject", headers=ADM, json={"reason": "Not something we process"})
        st = await status(sid)
        check("reject -> Rejected with the reviewer's reason, nothing re-queued", r.status_code == 200 and st["status"] == "Rejected" and any("Not something we process" in x for x in st["disposition_reasons"]) and len(await jobs(sid)) == 1, f"{r.text[:80]} {st['disposition_reasons']}")
        check("rejected upload keeps its file, listed as kind other", await file_exists(sid) and any(x["source_document_id"] == sid and x["kind"] == "other" for x in (await c.get("/invoices/rejected", params={"kind": "other"}, headers=ADM)).json()))
        r = await c.post(f"/documents/type-reviews/{rid}/choose", headers=ADM, json={"doc_type": "invoice"})
        check("choose after reject -> 409", r.status_code == 409)

        sid, used = await up("f_po_as_inv.docx")
        st = await status(sid)
        r = await c.post(f"/documents/type-reviews/{st['review_id']}/choose", headers=ADM, json={"doc_type": "invoice"})
        await wait(sid, 2)
        st = await status(sid)
        check("PO wrongly picked as invoice -> rejected by invoice validation, all reasons", st["status"] == "Rejected" and len(st["disposition_reasons"]) >= 2, f"{st['status']} {st['disposition_reasons']}")
        check("...no invoice row, file kept", (await q("select count(*) from invoices where source_document_id=:i", i=sid))[0][0] == 0 and await file_exists(sid))

        sid, used = await up("f_nonum_po.docx")
        st = await status(sid)
        r = await c.post(f"/documents/type-reviews/{st['review_id']}/choose", headers=ADM, json={"doc_type": "po"})
        await wait(sid, 2)
        st = await status(sid)
        check("picked PO, no PO number -> compulsory-number rule rejects it", st["status"] == "Rejected" and any("PO number" in x for x in st["disposition_reasons"]), f"{st['status']} {st['disposition_reasons']}")

        # ===================================================== boundary and failures
        sid, used = await up("f_edge75.docx")
        st = await status(sid)
        check("confidence exactly 75 is enough (inclusive) -> invoice path", st["status"] == "Invoice created", st["status"])
        sid, used = await up("f_choose_twice.docx")
        st = await status(sid)
        check("confidence 74 -> review", st["status"] == "Needs review" and st["review_kind"] == "document_type", st["status"])
        rid = st["review_id"]
        rs = await asyncio.gather(
            c.post(f"/documents/type-reviews/{rid}/choose", headers=ADM, json={"doc_type": "invoice"}),
            c.post(f"/documents/type-reviews/{rid}/choose", headers=ADM, json={"doc_type": "invoice"}),
        )
        check("two simultaneous decisions: one 200, one 409", sorted(x.status_code for x in rs) == [200, 409], str([x.status_code for x in rs]))
        await wait(sid, 2)
        await asyncio.sleep(2)
        check("...and the document is processed exactly once more", len(await jobs(sid)) == 2 and (await q("select count(*) from invoices where invoice_number='INV-IN-008'"))[0][0] == 1)

        sid, used = await up("f_boom.docx")
        st = await status(sid)
        check("classifier unavailable -> upload Failed with the error, no business rows", st["status"] == "Failed" and "model unavailable" in (st["error"] or "") and (await q("select count(*) from invoices where source_document_id=:i", i=sid))[0][0] == 0, f"{st['status']} {st['error']}")

        sid, used = await up("f_confident_other.docx")
        st = await status(sid)
        check("a confident (even if wrong) 'other' is rejected and the file is kept", st["status"] == "Rejected" and await file_exists(sid), st["status"])

        # ===================================================== ownership
        sid, used = await up("f_fa.docx", H=FA)
        st = await status(sid, H=FA)
        rid = st["review_id"]
        check("FA's doubtful upload is in FA's queue only", rid in [x["review_id"] for x in (await c.get("/documents/type-reviews", headers=FA)).json()] and rid not in [x["review_id"] for x in (await c.get("/documents/type-reviews", headers=ADM)).json()])
        check("another user cannot see, choose or reject it (403)", (await c.get(f"/documents/type-reviews/{rid}", headers=ADM)).status_code == 403 and (await c.post(f"/documents/type-reviews/{rid}/choose", headers=ADM, json={"doc_type": "invoice"})).status_code == 403 and (await c.post(f"/documents/type-reviews/{rid}/reject", headers=ADM, json={"reason": "x"})).status_code == 403)
        r = await c.post(f"/documents/type-reviews/{rid}/choose", headers=FA, json={"doc_type": "invoice"})
        await wait(sid, 2)
        check("the uploader can decide", r.status_code == 200 and (await status(sid, H=FA))["status"] == "Invoice created")
        check("unknown review id -> 404", (await c.get("/documents/type-reviews/00000000-0000-0000-0000-000000000000", headers=ADM)).status_code == 404)
        check("no token -> 401", (await c.get("/documents/type-reviews")).status_code == 401)

        # ===================================================== final state
        n_inv = (await q("select count(*) from invoices"))[0][0]
        n_po = (await q("select count(*) from purchase_orders"))[0][0]
        check("final: invoices INV-IN-001/002/008/009/007, POs PO-IN-1/PO-IN-2", n_inv == 5 and n_po == 2, f"invoices={n_inv} pos={n_po}")
        stuck = await q("select count(*) from jobs where state not in ('Succeeded','Failed','DeadLettered')")
        check("no job left running", stuck[0][0] == 0)
    await worker.stop(grace_seconds=5)
    await engine.dispose()
    print(f"RESULT {sum(results)}/{len(results)} passed")


asyncio.run(main())
