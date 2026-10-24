import asyncio
import sys
import tempfile
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

sys.path.insert(0, "src")
import httpx
from docx import Document
from sqlalchemy import text

import core.services.invoices.fx as invoice_fx
from api.rest.app import app
from config.settings import get_settings
from core.services.fx.ecb import FxRate, FxServiceError
from data.clients.db import SessionLocal, engine

s = get_settings()
assert "ar_fx2" in s.database_url, "refusing to run against a non-test database"
OURS = "33AAHCG2210A1Z7"
results = []
calls = []  # every rate lookup the pipeline makes
real_fetch = invoice_fx.fetch_rate_to_inr


async def counting_fetch(currency, on_date):
    calls.append((currency, str(on_date)))
    return await real_fetch(currency, on_date)


invoice_fx.fetch_rate_to_inr = counting_fetch


def check(name, cond, extra=""):
    results.append(bool(cond))
    print(f"RESULT {'PASS' if cond else 'FAIL'} {name} {extra}"[:320])


def make_invoice(path, *, number, currency, total, issued="05/10/2026", to="Zenith Softlabs Pvt Ltd",
                 gstin="33AABCZ1234F1Z5", title="TAX INVOICE"):
    d = Document()
    d.add_heading(title, 1)
    d.add_paragraph(f"Genworx Technologies Pvt Ltd\n12 Anna Salai, Chennai 600002\nGSTIN: {OURS}")
    d.add_paragraph(f"Invoice No: {number}\nInvoice Date: {issued}\nDue Date: 30/11/2026")
    d.add_paragraph(f"Bill To:\n{to}\nBengaluru\nGSTIN: {gstin}")
    t = d.add_table(rows=2, cols=3)
    t.style = "Table Grid"
    for i, h in enumerate(("Description", "Qty", f"Amount ({currency})")):
        t.rows[0].cells[i].text = h
    for i, v in enumerate(("Consulting services", "1", total)):
        t.rows[1].cells[i].text = v
    d.add_paragraph(f"Total {currency} {total}")
    d.save(path)


async def main():
    import testsupport
    worker = await testsupport.start_worker()
    tmp = Path(tempfile.mkdtemp())
    make_invoice(tmp / "inr.docx", number="FX-INR", currency="INR", total="100000.00")
    make_invoice(tmp / "usd.docx", number="FX-USD", currency="USD", total="1000.00", issued="05/10/2026")
    make_invoice(tmp / "eur_sun.docx", number="FX-EUR", currency="EUR", total="500.00", issued="11/10/2026")
    make_invoice(tmp / "qar.docx", number="FX-QAR", currency="QAR", total="2000.00")
    make_invoice(tmp / "usd_down.docx", number="FX-DOWN", currency="USD", total="300.00", issued="06/10/2026")
    make_invoice(tmp / "usd_stale.docx", number="FX-STALE", currency="USD", total="400.00", issued="07/10/2026")
    make_invoice(tmp / "usd_unknown.docx", number="FX-UNK", currency="USD", total="100.00", to="Nobody Ltd", gstin="07AAACN1111A1Z1")
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t", timeout=900) as c:
        testsupport.install(c)
        tok = (await c.post("/auth/login", json={"email": s.seed_admin_email, "password": s.seed_admin_password})).json()["access_token"]
        H = {"Authorization": f"Bearer {tok}"}
        cust = lambda **kw: {"contact": "Contact", "country": "India", "billing_currency": "INR", **kw}
        assert (await c.post("/customers", headers=H, json=cust(customer_name="Zenith Softlabs Pvt Ltd", email="ap@zenith.com", tax_id="33AABCZ1234F1Z5", deducts_tds=True))).status_code == 201

        async def up(name):
            r = await c.post("/documents/upload", headers=H, files={"file": (name, (tmp / name).read_bytes())})
            assert r.status_code == 202, r.text
            sid = r.json()["source_document_id"]
            for _ in range(300):
                row = next(u for u in (await c.get("/documents/uploads", headers=H)).json() if u["source_document_id"] == sid)
                if row["processing_status"] in ("Succeeded", "Failed"):
                    break
                await asyncio.sleep(1)
            return row, (await c.get(f"/documents/uploads/{sid}", headers=H)).json()

        async def endpoint_rate(cur, d):
            return (await c.get("/fx/rate", params={"currency": cur, "date": d}, headers=H)).json()

        # 1. INR: the normal flow, and the rate service is never touched
        calls.clear()
        row, det = await up("inr.docx")
        check("INR invoice -> created", row["disposition"] == "invoice_created", str(row["disposition"]))
        check("INR invoice made NO rate lookup", calls == [], str(calls))
        async with SessionLocal() as db:
            r = (await db.execute(text("select invoice_currency, amount_inr, amount_foreign, fx_rate_invoice, tds_amount, collectible_total from invoices where invoice_number='FX-INR'"))).one()
        check("INR invoice stored as before: rate 1, no foreign amount, TDS 10%", (r[0].strip(), r[1], r[2], r[3], r[4], r[5]) == ("INR", Decimal("100000.00"), None, Decimal("1.00000000"), Decimal("10000.00"), Decimal("90000.00")), str(tuple(r)))

        # 2. USD on a weekday
        calls.clear()
        row, det = await up("usd.docx")
        check("USD invoice -> created (no longer sent to review)", row["disposition"] == "invoice_created", str(det["disposition_reasons"]))
        check("exactly one rate lookup, for USD on the invoice date", calls == [("USD", "2026-10-05")], str(calls))
        ref = await endpoint_rate("USD", "2026-10-05")
        rate = Decimal(ref["inr_per_unit"])
        async with SessionLocal() as db:
            r = (await db.execute(text("select invoice_currency, amount_foreign, fx_rate_invoice, amount_inr, tds_amount, collectible_total from invoices where invoice_number='FX-USD'"))).one()
        exp_inr = (Decimal("1000") * rate).quantize(Decimal("0.01"), ROUND_HALF_UP)
        exp_tds = (exp_inr * Decimal("0.10")).quantize(Decimal("0.01"), ROUND_HALF_UP)
        check("USD invoice: currency, foreign amount, frozen ECB rate", (r[0].strip(), r[1], r[2]) == ("USD", Decimal("1000.00"), rate), str(tuple(r)))
        check("USD invoice: INR amount = 1000 x rate; TDS = 10% of the INR amount; collectible = INR - TDS", (r[3], r[4], r[5]) == (exp_inr, exp_tds, exp_inr - exp_tds), f"{tuple(r)} vs {exp_inr}/{exp_tds}")
        route = next(x for x in det["steps"] if x["step_name"] == "route")["output_summary"]
        check("route step records the rate used", route["fx"] and route["fx"]["currency"] == "USD" and route["fx"]["inr_per_unit"] == str(rate) and route["fx"]["days_back"] == 0, str(route.get("fx")))
        async with SessionLocal() as db:
            reason = (await db.execute(text("select reason from state_log where entity_type='invoice' and reason like '%converted%'"))).scalar()
        check("invoice timeline mentions the conversion", reason and "USD converted at" in reason, str(reason))

        # 3. EUR dated on a Sunday -> Friday's rate
        calls.clear()
        row, det = await up("eur_sun.docx")
        fri = await endpoint_rate("EUR", "2026-10-09")
        async with SessionLocal() as db:
            r = (await db.execute(text("select fx_rate_invoice, amount_inr from invoices where invoice_number='FX-EUR'"))).one()
        check("EUR invoice dated Sunday 2026-10-11 -> uses Friday's rate", row["disposition"] == "invoice_created" and r[0] == Decimal(fri["inr_per_unit"]) and r[1] == (Decimal("500") * Decimal(fri["inr_per_unit"])).quantize(Decimal("0.01"), ROUND_HALF_UP), str(tuple(r)))
        route = next(x for x in det["steps"] if x["step_name"] == "route")["output_summary"]
        check("... and the step log says it was 2 days back", route["fx"]["rate_date"] == "2026-10-09" and route["fx"]["days_back"] == 2, str(route["fx"]))

        # 4. a currency the ECB does not publish
        row, det = await up("qar.docx")
        check("QAR invoice -> review, reason says the ECB may not publish it and a rate can be entered", row["disposition"] == "review" and any("QAR" in x and "ECB" in x and "enter the INR rate" in x for x in det["disposition_reasons"]), str(det["disposition_reasons"]))
        check("... and no invoice was created", row["invoice_id"] is None)
        r = await c.post(f"/invoices/reviews/{row['review_id']}/approve", headers=H, json={"reason": "try"})
        check("approving the QAR invoice without a rate -> 422 (still no rate), stays pending", r.status_code == 422 and "No ECB rate" in r.text, r.text[:140])
        r = await c.post(f"/invoices/reviews/{row['review_id']}/approve", headers=H, json={"reason": "rate from the bank", "fx_rate": "-1"})
        check("a rate that is not positive -> 422", r.status_code == 422, r.text[:140])
        r = await c.post(f"/invoices/reviews/{row['review_id']}/approve", headers=H, json={"reason": "rate from the bank", "fx_rate": "22.75"})
        check("approving with a typed rate -> 200 and the invoice is created", r.status_code == 200 and r.json()["invoice_id"], r.text[:140])
        async with SessionLocal() as db:
            r = (await db.execute(text("select i.fx_rate_invoice, i.amount_inr, (select reason from state_log where entity_type='invoice' and entity_id=i.id order by occurred_at limit 1) from invoices i where invoice_number='FX-QAR'"))).one()
        check("... frozen at the typed rate, and the timeline says it was entered by a person", r[0] == Decimal("22.75") and r[1] == Decimal("45500.00") and "entered by a person" in r[2], str(tuple(r)))

        # 5. ECB down at upload -> review; later approval (service back) creates the invoice
        async def down(currency, on_date):
            calls.append((currency, str(on_date)))
            raise FxServiceError("down")

        invoice_fx.fetch_rate_to_inr = down
        row, det = await up("usd_down.docx")
        check("ECB unreachable -> review (not failed, not lost)", row["disposition"] == "review" and row["processing_status"] == "Succeeded" and any("not available right now" in x for x in det["disposition_reasons"]), str(det["disposition_reasons"]))
        rid = row["review_id"]
        r = await c.post(f"/invoices/reviews/{rid}/approve", headers=H, json={"reason": "retry while ECB is down"})
        check("approving while the ECB is still down -> 502, stays pending", r.status_code == 502 and (await c.get(f"/invoices/reviews/{rid}", headers=H)).json()["status"] == "Pending", r.text[:120])
        invoice_fx.fetch_rate_to_inr = counting_fetch  # service is back
        r = await c.post(f"/invoices/reviews/{rid}/approve", headers=H, json={"reason": "ECB is back, rate fetched on approval"})
        check("approving after it is back -> invoice created", r.status_code == 200 and r.json()["invoice_id"], r.text[:120])
        ref = await endpoint_rate("USD", "2026-10-06")
        async with SessionLocal() as db:
            r = (await db.execute(text("select fx_rate_invoice, amount_inr, created_by is not null from invoices where invoice_number='FX-DOWN'"))).one()
        check("... stored with the ECB rate of the invoice date", r[0] == Decimal(ref["inr_per_unit"]) and r[1] == (Decimal("300") * Decimal(ref["inr_per_unit"])).quantize(Decimal("0.01"), ROUND_HALF_UP), str(tuple(r)))

        # 6. stale rate -> review
        async def stale(currency, on_date):
            calls.append((currency, str(on_date)))
            return FxRate(currency, on_date, date_minus(on_date, 7), Decimal("90"), Decimal("100"), Decimal("1.1"))

        invoice_fx.fetch_rate_to_inr = stale
        row, det = await up("usd_stale.docx")
        check("rate more than 5 days older than the invoice date -> review 'too old'", row["disposition"] == "review" and any("too old" in x for x in det["disposition_reasons"]), str(det["disposition_reasons"]))
        invoice_fx.fetch_rate_to_inr = counting_fetch

        # 7. unknown customer + foreign currency: no wasted rate lookup
        calls.clear()
        row, det = await up("usd_unknown.docx")
        check("USD invoice for an unknown customer -> review, with NO rate lookup", row["disposition"] == "review" and calls == [], f"{row['disposition']} {calls}")

        async with SessionLocal() as db:
            n = (await db.execute(text("select count(*) from invoices"))).scalar()
            check("invoices created: INR, USD, EUR, the QAR one approved with a typed rate, and the one approved after the ECB came back = 5", n == 5, str(n))
    await worker.stop(grace_seconds=5)
    await engine.dispose()
    print(f"RESULT {sum(results)}/{len(results)} passed")


from datetime import timedelta


def date_minus(d, n):
    return d - timedelta(days=n)


asyncio.run(main())
