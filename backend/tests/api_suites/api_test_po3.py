import asyncio
import sys
import tempfile
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, "src")
import httpx
from docx import Document
from sqlalchemy import text

from api.rest.app import app
from config.settings import get_settings
from data.clients.db import SessionLocal, engine

s = get_settings()
assert "ar_po3" in s.database_url, "refusing to run against a non-test database"
OURS = "33AAHCG2210A1Z7"
results = []


def check(name, cond, extra=""):
    results.append(bool(cond))
    print(f"RESULT {'PASS' if cond else 'FAIL'} {name} {extra}"[:300])


def make_po(path, *, number, buyer="Zenith Softlabs Pvt Ltd", gstin="33AABCZ1234F1Z5", lines, total):
    d = Document()
    d.add_heading("PURCHASE ORDER", 1)
    d.add_paragraph(f"{buyer}\nBengaluru\nGSTIN: {gstin}")
    d.add_paragraph(f"PO No: {number}\nPO Date: 02/10/2026")
    d.add_paragraph(f"Supplier: Genworx Technologies Pvt Ltd, GSTIN {OURS}")
    t = d.add_table(rows=1 + len(lines), cols=3)
    t.style = "Table Grid"
    for i, h in enumerate(("Description", "Qty", "Unit price (INR)")):
        t.rows[0].cells[i].text = h
    for r, row in enumerate(lines, start=1):
        for i, v in enumerate(row):
            t.rows[r].cells[i].text = v
    d.add_paragraph(f"Order total INR {total}")
    d.save(path)


def make_invoice(path, *, number, po, lines, total, gstin="33AABCZ1234F1Z5", to="Zenith Softlabs Pvt Ltd"):
    d = Document()
    d.add_heading("TAX INVOICE", 1)
    d.add_paragraph(f"Genworx Technologies Pvt Ltd\n12 Anna Salai, Chennai 600002\nGSTIN: {OURS}")
    d.add_paragraph(f"Invoice No: {number}\nInvoice Date: 05/10/2026\nDue Date: 04/11/2026")
    if po:
        d.add_paragraph(f"Customer PO No: {po}")
    d.add_paragraph(f"Bill To:\n{to}\nBengaluru\nGSTIN: {gstin}")
    t = d.add_table(rows=1 + len(lines), cols=4)
    t.style = "Table Grid"
    for i, h in enumerate(("Description", "Qty", "Unit price", "Amount (INR)")):
        t.rows[0].cells[i].text = h
    for r, row in enumerate(lines, start=1):
        for i, v in enumerate(row):
            t.rows[r].cells[i].text = v
    d.add_paragraph(f"Total INR {total}")
    d.save(path)


async def main():
    import testsupport
    worker = await testsupport.start_worker()
    tmp = Path(tempfile.mkdtemp())
    make_po(tmp / "po5000.docx", number="PO-5000", lines=(("Laptop", "10", "50000.00"), ("Monitor", "5", "12000.00")), total="560000.00")
    make_po(tmp / "po6000.docx", number="PO-6000", buyer="Meridian Tech Solutions Pvt Ltd", gstin="29AABCM5678K1Z3", lines=(("Server", "2", "100000.00"),), total="200000.00")
    make_po(tmp / "po8000.docx", number="PO-8000", lines=(("Laptop", "5", "50000.00"),), total="250000.00")
    make_invoice(tmp / "i1.docx", number="INV-3001", po="PO-5000", lines=(("Laptop", "6", "50000.00", "300000.00"),), total="300000.00")
    make_invoice(tmp / "i2.docx", number="INV-3002", po="PO-5000", lines=(("Laptop", "6", "50000.00", "300000.00"),), total="300000.00")
    make_invoice(tmp / "i3.docx", number="INV-3003", po="PO-5000", lines=(("Keyboard", "1", "3000.00", "3000.00"),), total="3000.00")
    make_invoice(tmp / "i4.docx", number="INV-3004", po="PO-7777", lines=(("Cables", "3", "100.00", "300.00"),), total="300.00")
    make_invoice(tmp / "i5.docx", number="INV-3005", po="PO-6000", lines=(("Server", "1", "100000.00", "100000.00"),), total="100000.00")
    make_invoice(tmp / "i6.docx", number="INV-3006", po=None, lines=(("Training", "2", "1000.00", "2000.00"),), total="2000.00")
    make_invoice(tmp / "i7.docx", number="INV-3007", po="PO-8000", lines=(("Laptop", "2", "50000.00", "100000.00"),), total="100000.00")

    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t", timeout=900) as c:
        testsupport.install(c)
        tok = (await c.post("/auth/login", json={"email": s.seed_admin_email, "password": s.seed_admin_password})).json()["access_token"]
        H = {"Authorization": f"Bearer {tok}"}
        cust = lambda **kw: {"contact": "Contact", "country": "India", "billing_currency": "INR", **kw}
        for body in (
            cust(customer_name="Zenith Softlabs Pvt Ltd", email="ap@zenith.com", gstin="33AABCZ1234F1Z5", deducts_tds=True),
            cust(customer_name="Meridian Tech Solutions Pvt Ltd", email="ap@meridian.com", gstin="29AABCM5678K1Z3", deducts_tds=False),
        ):
            assert (await c.post("/customers", headers=H, json=body)).status_code == 201

        async def up(kind, name):
            ep = "/documents"
            r = await c.post(f"{ep}/upload", headers=H, files={"file": (name, (tmp / name).read_bytes())})
            assert r.status_code == 202, (name, r.text)
            sid = r.json()["source_document_id"]
            for _ in range(300):
                row = next(u for u in (await c.get(f"{ep}/uploads", headers=H)).json() if u["source_document_id"] == sid)
                if row["processing_status"] in ("Succeeded", "Failed"):
                    break
                await asyncio.sleep(1)
            detail = (await c.get(f"{ep}/uploads/{sid}", headers=H)).json()
            return row, detail

        row, det = await up("po", "po5000.docx")
        check("PO-5000 created", row["disposition"] == "po_created", str(row["disposition"]))
        row, det = await up("po", "po6000.docx")
        check("PO-6000 (Meridian) created", row["disposition"] == "po_created", str(row["disposition"]))

        # PO-first
        r1, _ = await up("inv", "i1.docx")
        check("I1 (6 of 10 laptops) -> created, not held", r1["disposition"] == "invoice_created" and r1["po_hold"] is False, str(r1))
        r2, d2 = await up("inv", "i2.docx")
        check("I2 (6 more laptops, 12 > 10) -> created but ON HOLD", r2["disposition"] == "invoice_created" and r2["po_hold"] is True, str(r2["po_hold"]))
        r3, _ = await up("inv", "i3.docx")
        check("I3 (item not on the PO) -> created, on hold", r3["disposition"] == "invoice_created" and r3["po_hold"] is True)
        r4, _ = await up("inv", "i4.docx")
        check("I4 (PO-7777 exists nowhere) -> created, unlinked, not held", r4["disposition"] == "invoice_created" and r4["po_hold"] is False)
        r6, _ = await up("inv", "i6.docx")
        check("I6 (no PO number) -> created, not held", r6["disposition"] == "invoice_created" and r6["po_hold"] is False)

        # G2
        r5, d5 = await up("inv", "i5.docx")
        check("I5 quotes Meridian's PO-6000 for Zenith -> REVIEW naming the other customer", r5["disposition"] == "review" and "Meridian" in " ".join(d5["disposition_reasons"]) and "PO-6000" in " ".join(d5["disposition_reasons"]), str(d5["disposition_reasons"]))
        resp = await c.post(f"/invoices/reviews/{r5['review_id']}/approve", headers=H, json={"reason": "Customer confirmed the PO number is correct"})
        check("approving I5 in review creates the invoice, unlinked and not held", resp.status_code == 200 and resp.json()["invoice_id"] and resp.json()["po_hold"] is False, resp.text[:150])

        # invoice-first
        r7, _ = await up("inv", "i7.docx")
        check("I7 (PO-8000 not uploaded yet) -> created, unlinked", r7["disposition"] == "invoice_created" and r7["po_hold"] is False)
        row, det = await up("po", "po8000.docx")
        check("PO-8000 arrives: created and retro-links INV-3007, which passes", row["disposition"] == "po_created" and [(x["invoice_number"], x["passed"]) for x in det["linked_invoices"]] == [("INV-3007", True)], str(det["linked_invoices"]))

        async with SessionLocal() as db:
            q = lambda sql, **p: db.execute(text(sql), p)
            links = {r[0]: (r[1], r[2]) for r in (await q("select i.invoice_number, p.po_number, m.match_direction from po_invoice_matches m join invoices i on i.id=m.invoice_id join purchase_orders p on p.id=m.po_id")).all()}
            check("links: 3001-3003 po_first to PO-5000, 3007 invoice_first to PO-8000; no others",
                  links == {"INV-3001": ("PO-5000", "po_first"), "INV-3002": ("PO-5000", "po_first"), "INV-3003": ("PO-5000", "po_first"), "INV-3007": ("PO-8000", "invoice_first")}, str(links))
            ledger = {r[0]: r[1] for r in (await q("select i.invoice_number, sum(f.quantity_invoiced) from po_fulfillments f join invoices i on i.id=f.invoice_id group by 1 order by 1")).all()}
            check("ledger only for passing invoices (3001: 6, 3007: 2)", ledger == {"INV-3001": Decimal("6.0000"), "INV-3007": Decimal("2.0000")}, str(ledger))
            holds = {r[0] for r in (await q("select invoice_number from invoices where po_hold")).all()}
            check("po_hold exactly INV-3002 and INV-3003", holds == {"INV-3002", "INV-3003"}, str(holds))
            pend = {r[0] for r in (await q("select i.invoice_number from pending_invoice_changes p join invoices i on i.id=p.invoice_id where p.reason='po_mismatch' and p.status='Pending'")).all()}
            check("po_mismatch Pending rows exactly for those two", pend == {"INV-3002", "INV-3003"}, str(pend))
            check("hold is logged on the invoice timeline", (await q("select count(*) from state_log where entity_type='invoice' and reason like 'On PO hold%'")).scalar() == 2)
            nop = (await q("select count(*) from invoices where invoice_number in ('INV-3004','INV-3005','INV-3006')")).scalar()
            check("INV-3004/3005/3006 exist and are not linked", nop == 3 and not ({"INV-3004", "INV-3005", "INV-3006"} & set(links)))
            rev_ok = (await q("select count(*) from pending_invoice_changes where reason='extraction_review' and status='Approved'")).scalar()
            check("I5 extraction review recorded as Approved", rev_ok == 1)
        # detail of a held upload still shows invoice_created and the hold
        det2 = (await c.get(f"/documents/uploads/{r2['source_document_id']}", headers=H)).json()
        check("held invoice's upload detail: invoice_created + po_hold", det2["disposition"] == "invoice_created" and det2["po_hold"] is True)
    await worker.stop(grace_seconds=5)
    await engine.dispose()
    print(f"RESULT {sum(results)}/{len(results)} passed")


asyncio.run(main())
