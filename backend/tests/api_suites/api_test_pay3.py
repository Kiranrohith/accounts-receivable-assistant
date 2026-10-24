"""P3: a remittance advice uploaded through POST /documents/upload becomes a payment AND is applied.

Real worker, real LLM (classify + extract), real database. Throwaway DB only.
"""

import asyncio
import sys
import tempfile
from datetime import date
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, "src")
import httpx
from docx import Document
from sqlalchemy import text

from api.rest.app import app
from config.settings import get_settings
from core.services.payments.identities import verify_money_identities
from data.clients.db import SessionLocal, engine
from data.models.postgres import Invoice

s = get_settings()
assert "ar_pay3" in s.database_url, "refusing to run against a non-test database"
results = []


def check(name, cond, extra=""):
    results.append(bool(cond))
    print(f"RESULT {'PASS' if cond else 'FAIL'} {name} {extra}"[:320])


def make_advice(path, *, payer="Zenith Softlabs Pvt Ltd", from_email=None, ref="UTR2026101500123",
                paid_on="15/10/2026", amount="900.00", currency="INR", invoices=(("ZEN-INV-101", "900.00"),),
                title="PAYMENT ADVICE", bank="HDFC Bank"):
    d = Document()
    d.add_heading(title, 1)
    d.add_paragraph(f"From: {payer}" + (f"\nEmail: {from_email}" if from_email else ""))
    d.add_paragraph("To: Genworx Technologies Pvt Ltd")
    d.add_paragraph(f"Paying bank: {bank}")
    if ref:
        d.add_paragraph(f"Payment reference (UTR): {ref}")
    d.add_paragraph(f"Payment date: {paid_on}")
    d.add_paragraph("Mode: NEFT")
    if invoices:
        t = d.add_table(rows=1 + len(invoices), cols=2)
        t.style = "Table Grid"
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Invoice No", f"Amount paid ({currency})"
        for r, (no, amt) in enumerate(invoices, start=1):
            t.rows[r].cells[0].text, t.rows[r].cells[1].text = no, amt
    if amount:
        d.add_paragraph(f"Total amount paid: {currency} {amount}")
    d.add_paragraph("Please acknowledge receipt of this payment.")
    d.save(path)


def make_challan(path):
    d = Document()
    d.add_heading("DELIVERY CHALLAN", 1)
    d.add_paragraph("Challan No: DC-77\nDate: 03/10/2026\nConsignee: Zenith Softlabs Pvt Ltd, Bengaluru")
    t = d.add_table(rows=2, cols=3)
    t.style = "Table Grid"
    for i, h in enumerate(("Item", "Qty", "Remarks")):
        t.rows[0].cells[i].text = h
    for i, v in enumerate(("Laptop bags", "20", "Delivered by road")):
        t.rows[1].cells[i].text = v
    d.add_paragraph("Goods delivered in good condition. Receiver's signature: ________")
    d.save(path)


async def seed_invoice(customer_id, number, collectible="900"):
    async with SessionLocal() as db:
        inv = Invoice(
            customer_id=customer_id, invoice_number=number, invoice_date=date(2026, 10, 1),
            due_date=date(2026, 10, 31), invoice_currency="INR", amount_inr=Decimal("1000"),
            collectible_total=Decimal(collectible),
        )
        db.add(inv)
        await db.commit()


async def main():
    import testsupport
    worker = await testsupport.start_worker()
    tmp = Path(tempfile.mkdtemp())
    files = {
        "adv_ok.docx": dict(from_email="accounts@zenith.com"),
        "adv_email_only.docx": dict(payer="M T Solutions", from_email="ap@meridian.com", ref="UTR2026101500200",
                                    invoices=(), amount="1111.00"),
        "adv_name_only.docx": dict(ref="UTR2026101500300", invoices=(), amount="2500.00"),
        "adv_unknown.docx": dict(payer="Nobody Traders Pvt Ltd", ref="UTR2026101500400", invoices=(), amount="700.00"),
        "adv_no_amount.docx": dict(ref="UTR2026101500500", amount=None, invoices=(("ZEN-INV-102", ""),)),
        "adv_conflict.docx": dict(payer="Zenith Softlabs Pvt Ltd", from_email="ap@meridian.com",
                                  ref="UTR2026101500600", invoices=(("ZEN-INV-102", "900.00"),)),
        "adv_dup.docx": dict(ref="UTR2026101500123", invoices=(("ZEN-INV-102", "900.00"),)),
        "adv_junk1.docx": dict(ref="NEFT", invoices=(("MER-INV-201", "450.00"),), payer="Meridian Tech Solutions Pvt Ltd", amount="450.00"),
        "adv_junk2.docx": dict(ref="NEFT", invoices=(("MER-INV-202", "450.00"),), payer="Meridian Tech Solutions Pvt Ltd", amount="450.00"),
        "adv_two_owners.docx": dict(ref="UTR2026101500700", invoices=(("ZEN-INV-102", "400.00"), ("MER-INV-201", "500.00")), amount="900.00"),
        "adv_usd.docx": dict(ref="SWIFT2026101500800", invoices=(), amount="1200.00", currency="USD"),
        "adv_short.docx": dict(ref="UTR2026101500900", invoices=(("ZEN-INV-102", "600.00"),), amount="600.00"),
        "adv_split.docx": dict(ref="UTR2026101501000", payer="Meridian Tech Solutions Pvt Ltd", amount="800.00",
                               invoices=(("MER-INV-204", "300.00"), ("MER-INV-205", "500.00"))),
        "adv_over.docx": dict(ref="UTR2026101501100", payer="Meridian Tech Solutions Pvt Ltd", amount="1200.00",
                              invoices=(("MER-INV-203", "1200.00"),)),
    }
    for name, kw in files.items():
        make_advice(tmp / name, **kw)
    make_challan(tmp / "challan.docx")

    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t", timeout=900) as c:
        testsupport.install(c)
        tok = (await c.post("/auth/login", json={"email": s.seed_admin_email, "password": s.seed_admin_password})).json()["access_token"]
        H = {"Authorization": f"Bearer {tok}"}
        cust = lambda **kw: {"contact": "Contact", "country": "India", "billing_currency": "INR", **kw}
        for body in (
            cust(customer_name="Zenith Softlabs Pvt Ltd", email="ap@zenith.com", tax_id="33AABCZ1234F1Z5", deducts_tds=True),
            cust(customer_name="Meridian Tech Solutions Pvt Ltd", email="ap@meridian.com", tax_id="29AABCM5678K1Z3", deducts_tds=False),
        ):
            assert (await c.post("/customers", headers=H, json=body)).status_code == 201
        async with SessionLocal() as db:
            cid = {r[0]: r[1] for r in (await db.execute(text("select customer_name, id from customers"))).all()}
        ZEN, MER = cid["Zenith Softlabs Pvt Ltd"], cid["Meridian Tech Solutions Pvt Ltd"]
        for number in ("ZEN-INV-101", "ZEN-INV-102"):
            await seed_invoice(ZEN, number)
        for number in ("MER-INV-201", "MER-INV-202"):
            await seed_invoice(MER, number, "450")
        await seed_invoice(MER, "MER-INV-203", "1000")
        await seed_invoice(MER, "MER-INV-204", "300")
        await seed_invoice(MER, "MER-INV-205", "500")

        ids = {}

        async def up(name):
            r = await c.post("/documents/upload", headers=H, files={"file": (name, (tmp / name).read_bytes())})
            assert r.status_code == 202, (name, r.text)
            ids[name] = r.json()["source_document_id"]
            return (await c.get(f"/documents/uploads/{ids[name]}", headers=H)).json()

        async def payment_of(name):
            async with SessionLocal() as db:
                row = (await db.execute(text(
                    "select p.id, p.state::text, p.customer_id, p.matched_layer::text, p.amount_received_inr, "
                    "p.amount_received_foreign, p.payment_currency, p.reference_number, p.unapplied_amount, "
                    "p.held_amount, p.credited_amount, p.duplicate_of, p.review_reasons, p.extracted_invoice_numbers, "
                    "p.payment_date from payments p where p.source_document_id=:d"), {"d": ids[name]})).first()
            return row

        # 1. a clean advice quoting a known invoice: L1, ready for the matcher
        u = await up("adv_ok.docx")
        p = await payment_of("adv_ok.docx")
        check("adv_ok: kind payment, status 'Payment applied'", u["document_kind"] == "payment" and u["status"] == "Payment applied", f"{u['document_kind']} {u['status']} {u.get('error')}")
        check("adv_ok: disposition payment_applied, payment linked", u["disposition"] == "payment_applied" and u["payment_id"] == str(p[0]) if p else False)
        check("adv_ok: payment Applied, customer Zenith, layer L1", p and p[1] == "Applied" and p[2] == ZEN and p[3] == "L1", str(p and p[1:4]))
        async with SessionLocal() as db:
            i101 = (await db.execute(text("select status::text, amount_paid from invoices where invoice_number='ZEN-INV-101'"))).first()
        check("adv_ok: invoice ZEN-INV-101 is now Paid, 900", i101 == ("Paid", Decimal("900.00")), str(i101))
        check("adv_ok: amount 900 INR, ref and date read", p and p[4] == Decimal("900.00") and p[5] is None and p[6] == "INR" and p[7] == "UTR2026101500123" and p[14] == date(2026, 10, 15), str(p and (p[4], p[6], p[7], p[14])))
        check("adv_ok: buckets zero", p and p[8] == 0 and p[9] == 0 and p[10] == 0)
        check("adv_ok: quoted invoice stored", p and any(i["invoice_number"] == "ZEN-INV-101" for i in (p[13] or [])), str(p and p[13]))
        det = (await c.get(f"/documents/uploads/{ids['adv_ok.docx']}", headers=H)).json()
        check("adv_ok: steps parse, classify, extract, validate, identify, route, match", [x["step_name"] for x in det["steps"]] == ["parse", "classify", "extract", "validate", "identify", "route", "match"], str([x["step_name"] for x in det["steps"]]))
        check("adv_ok: match step shows the allocation", det["steps"][-1]["output_summary"]["status"] == "applied" and det["steps"][-1]["output_summary"]["allocations"][0]["invoice_number"] == "ZEN-INV-101", str(det["steps"][-1]["output_summary"]))
        check("adv_ok: classification says payment", (det["classification"] or {}).get("doc_class") == "payment", str(det["classification"]))
        check("adv_ok: extraction shown, amount and reference visible", det["extraction"] and det["reference_number"] == "UTR2026101500123" and det["total_amount"] is not None and Decimal(str(det["total_amount"])) == Decimal("900"), f"{det['reference_number']} {det['total_amount']}")
        check("adv_ok: payer name shown as the customer name", det["customer_name"] and "Zenith" in det["customer_name"], str(det["customer_name"]))

        # 2. e-mail only -> L2
        await up("adv_email_only.docx")
        p = await payment_of("adv_email_only.docx")
        check("email only: identified as Meridian (L2) but no invoice combination fits 1111: Unapplied cash (2026-10-21 rule)", p and p[1] == "Unapplied" and p[2] == MER and p[3] == "L2", str(p and p[1:4]))

        # 3. name only -> L3
        await up("adv_name_only.docx")
        p = await payment_of("adv_name_only.docx")
        check("name only: identified as Zenith (L3) but nothing fits 2500: Unapplied cash (2026-10-21 rule)", p and p[1] == "Unapplied" and p[2] == ZEN and p[3] == "L3", str(p and p[1:4]))

        # 4. nobody we know -> review
        u = await up("adv_unknown.docx")
        p = await payment_of("adv_unknown.docx")
        check("unknown payer: NeedsReview, no customer, L4", p and p[1] == "NeedsReview" and p[2] is None and p[3] == "L4", str(p and p[1:4]))
        check("unknown payer: reason says it could not be identified", p and any("could not be identified" in r for r in (p[12] or [])), str(p and p[12]))
        check("unknown payer: status 'Needs review', review_kind payment, review_id is the payment", u["status"] == "Needs review" and u["review_kind"] == "payment" and u["review_id"] == str(p[0]), f"{u['status']} {u['review_kind']}")
        check("unknown payer: money buckets zero", p and p[8] == 0 and p[9] == 0 and p[10] == 0)
        d = (await c.get(f"/documents/uploads/{ids['adv_unknown.docx']}", headers=H)).json()
        check("unknown payer: reasons shown in the upload detail", any("could not be identified" in r for r in d["disposition_reasons"]), str(d["disposition_reasons"]))

        # 5. amount missing -> review, but the customer is still identified
        u = await up("adv_no_amount.docx")
        p = await payment_of("adv_no_amount.docx")
        check("no amount: NeedsReview, customer Zenith kept, amount null", p and p[1] == "NeedsReview" and p[2] == ZEN and p[4] is None, str(p and (p[1], p[2], p[4])))
        check("no amount: reason names the missing amount", p and any("does not state the amount" in r for r in (p[12] or [])), str(p and p[12]))

        # 6. quoted invoice belongs to Zenith but the sender is Meridian (C1)
        await up("adv_conflict.docx")
        p = await payment_of("adv_conflict.docx")
        check("conflict: NeedsReview with an identity conflict reason", p and p[1] == "NeedsReview" and any("Identity conflict" in r for r in (p[12] or [])), str(p and (p[1], p[12])))

        # 7. same UTR from the same customer as adv_ok -> duplicate
        u = await up("adv_dup.docx")
        p = await payment_of("adv_dup.docx")
        first = await payment_of("adv_ok.docx")
        check("duplicate reference: Rejected, duplicate_of the first", p and p[1] == "Rejected" and p[11] == first[0], str(p and (p[1], p[11])))
        check("duplicate: status 'Rejected' with a reason, buckets zero", u["status"] == "Rejected" and any("Duplicate" in r for r in (await c.get(f"/documents/uploads/{ids['adv_dup.docx']}", headers=H)).json()["disposition_reasons"]) and p[8] == 0 and p[10] == 0)

        # 8. junk references never dedup
        await up("adv_junk1.docx")
        await up("adv_junk2.docx")
        p1, p2 = await payment_of("adv_junk1.docx"), await payment_of("adv_junk2.docx")
        check("junk reference 'NEFT' twice: both stand and both apply, neither is a duplicate", p1 and p2 and p1[1] == "Applied" and p2[1] == "Applied" and p2[11] is None, str((p1 and p1[1], p2 and p2[1])))

        # 9. quoted invoices of two customers
        await up("adv_two_owners.docx")
        p = await payment_of("adv_two_owners.docx")
        check("two owners: NeedsReview 'different customers'", p and p[1] == "NeedsReview" and any("different customers" in r for r in (p[12] or [])), str(p and (p[1], p[12])))

        # 10. a foreign payment keeps its own currency until a match gives it a rate
        await up("adv_usd.docx")
        p = await payment_of("adv_usd.docx")
        check("USD payment (Zenith has no USD invoice): foreign amount kept, INR not set, review", p and p[1] == "NeedsReview" and p[6] == "USD" and p[5] == Decimal("1200.00") and p[4] is None, str(p and (p[1], p[4], p[5], p[6])))

        # 10b. a short payment from a quoted invoice: Partly paid
        u = await up("adv_short.docx")
        p = await payment_of("adv_short.docx")
        async with SessionLocal() as db:
            i102 = (await db.execute(text("select status::text, amount_paid, collectible_balance from invoices where invoice_number='ZEN-INV-102'"))).first()
        check("short payment 600 of 900: invoice PartlyPaid, balance 300", i102 == ("PartlyPaid", Decimal("600.00"), Decimal("300.00")), str(i102))
        check("short payment: status 'Payment applied', payment Applied", u["status"] == "Payment applied" and p[1] == "Applied", f"{u['status']} {p and p[1]}")

        # 10c. the advice states a split over two invoices
        await up("adv_split.docx")
        p = await payment_of("adv_split.docx")
        async with SessionLocal() as db:
            im = (await db.execute(text("select invoice_number, status::text from invoices where invoice_number in ('MER-INV-204','MER-INV-205') order by 1"))).all()
        check("stated split 300 + 500: both invoices Paid, payment Applied", p and p[1] == "Applied" and [x[1] for x in im] == ["Paid", "Paid"], str(p and p[1]) + str(im))

        # 10d. an overpayment: Paid, the excess becomes a credit note
        u = await up("adv_over.docx")
        p = await payment_of("adv_over.docx")
        async with SessionLocal() as db:
            i203 = (await db.execute(text("select status::text from invoices where invoice_number='MER-INV-203'"))).scalar()
            nts = (await db.execute(text("select note_number, amount, remaining_amount, status::text from notes where customer_id=:c"), {"c": MER})).all()
        check("overpayment 1200 on a 1000 invoice: Paid, payment Applied, 200 credited", i203 == "Paid" and p and p[1] == "Applied" and p[10] == Decimal("200.00"), f"{i203} {p and (p[1], p[10])}")
        check("overpayment: one credit note for 200, Open", len(nts) == 1 and nts[0][1] == Decimal("200.00") and nts[0][2] == Decimal("200.00") and nts[0][3] == "Open", str(nts))
        dd = (await c.get(f"/documents/uploads/{ids['adv_over.docx']}", headers=H)).json()
        check("overpayment: the match step names the credit note", dd["steps"][-1]["output_summary"]["credit_note"] == nts[0][0] if nts else False, str(dd["steps"][-1]["output_summary"]))

        # 11. an identical file is still refused
        r = await c.post("/documents/upload", headers=H, files={"file": ("again.docx", (tmp / "adv_ok.docx").read_bytes())})
        check("identical advice file -> 409", r.status_code == 409, str(r.status_code))

        # 12. the human type-review path: a challan rejected as 'other', sent back, chosen as a payment
        u = await up("challan.docx")
        check("challan: rejected as neither", u["status"] == "Rejected" and u["document_kind"] == "unknown", f"{u['status']} {u['document_kind']} {u.get('detected_type')}")
        rejected = [x for x in (await c.get("/invoices/rejected?kind=other", headers=H)).json() if x.get("source_document_id") == ids["challan.docx"]]
        if rejected:
            sb = await c.post(f"/invoices/rejected/{rejected[0]['rejected_id']}/send-to-review", headers=H)
            check("challan: sent back for a type decision", sb.status_code == 200, sb.text[:100])
            rid = sb.json()["review_id"]
            ch = await c.post(f"/documents/type-reviews/{rid}/choose", headers=H, json={"doc_type": "payment"})
            check("challan: choosing 'payment' is accepted", ch.status_code == 200 and ch.json()["processing"], ch.text[:120])
            await testsupport.wait_idle(ids["challan.docx"])
            d = (await c.get(f"/documents/uploads/{ids['challan.docx']}", headers=H)).json()
            async with SessionLocal() as db:
                jt = (await db.execute(text("select job_type from jobs where entity_id=:d order by created_at"), {"d": ids["challan.docx"]})).scalars().all()
            check("challan: a payment_extraction job ran, kind is now payment", "payment_extraction" in jt and d["document_kind"] == "payment", f"{jt} {d['document_kind']}")
            check("challan read as a payment advice has no usable payment: needs review", d["status"] in ("Needs review", "Rejected"), d["status"])
        else:
            check("challan: found in the rejected list", False, "not found")

        # 13. the whole book is still sound, and the list shows every payment kind
        async with SessionLocal() as db:
            found = await verify_money_identities(db)
            n = (await db.execute(text("select count(*) from payments"))).scalar()
            unique_ok = (await db.execute(text("select count(*) = count(distinct source_document_id) from payments"))).scalar()
        check("money identities hold", found == [], str([str(v) for v in found]))
        check("one payment per document", unique_ok, f"{n} payments")
        lst = (await c.get("/documents/uploads", headers=H)).json()
        check("list shows payment statuses", {"Payment applied", "Needs review", "Rejected"} <= {x["status"] for x in lst}, str({x["status"] for x in lst}))

    await worker.stop()
    await engine.dispose()
    print(f"RESULT {sum(results)}/{len(results)} passed")


asyncio.run(main())
