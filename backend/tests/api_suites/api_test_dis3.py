"""D3: an uploaded customer complaint becomes dispute(s). Real worker and database; the reader (LLM)
is replaced by scripted verdicts keyed by a marker in each document, so every route is deterministic."""

import asyncio
import re
import sys
import tempfile
from datetime import date
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, "src")
import httpx
from docx import Document
from sqlalchemy import text

import core.services.disputes.processing as dproc
import core.services.documents.intake as intake_mod
import core.services.payments.processing as pproc
from api.rest.app import app
from config.settings import get_settings
from core.services.disputes.extraction import DisputeOutcome
from core.services.documents.classification import ClassificationOutcome
from core.services.invoices.extraction import ExtractionError
from core.services.payments.extraction import PaymentOutcome
from core.services.payments.identities import verify_money_identities
from data.clients.db import SessionLocal, engine
from data.models.postgres import Customer, Invoice
from schemas.disputes.extraction import DisputeFieldConfidences, ExtractedDispute
from schemas.documents.classification import DocClassification, DocumentClass
from schemas.payments.extraction import ExtractedPayment, PaidInvoice, PaymentFieldConfidences

D = Decimal
s = get_settings()
assert "ar_dis3" in s.database_url, "refusing to run against a non-test database"
results = []


def check(name, cond, extra=""):
    results.append(bool(cond))
    print(f"RESULT {'PASS' if cond else 'FAIL'} {name} {extra}"[:330])


# ---- the scripted reader ------------------------------------------------------------------------
SCN: dict[str, dict] = {}
CALLS: dict[str, int] = {}


def marker(markdown: str) -> str:
    m = re.search(r"SCN-[A-Z0-9]+", markdown)
    assert m, markdown[:200]
    return m.group(0)


def cls(doc_class, confidence=95, also_payment=False, also_dispute=False):
    return DocClassification(doc_class=DocumentClass(doc_class), confidence=confidence, also_payment=also_payment,
                             also_dispute=also_dispute, reason="scripted")


def disp(*, invoices=(), name=None, email=None, gstin=None, category="Quality", cat_conf=92, related=False, summary="Goods arrived damaged"):
    return ExtractedDispute(customer_name=name, customer_gstin=gstin, from_email=email, invoice_numbers=list(invoices),
                            summary=summary, category=category, payment_related=related,
                            confidence=DisputeFieldConfidences(category=cat_conf if category else None, customer=90, invoices=90))


async def fake_classify(markdown):
    sc = SCN[marker(markdown)]
    CALLS[marker(markdown) + ":classify"] = CALLS.get(marker(markdown) + ":classify", 0) + 1
    return ClassificationOutcome(sc["cls"], "stub", 1)


async def fake_extract_dispute(markdown):
    key = marker(markdown)
    sc = SCN[key]
    CALLS[key + ":dispute"] = CALLS.get(key + ":dispute", 0) + 1
    if sc.get("fail_dispute_once") and CALLS[key + ":dispute"] == 1:
        raise ExtractionError("scripted: every model is down")
    return DisputeOutcome(sc["dispute"], "stub", 1)


async def fake_extract_payment(markdown):
    key = marker(markdown)
    sc = SCN[key]
    CALLS[key + ":payment"] = CALLS.get(key + ":payment", 0) + 1
    if sc.get("fail_payment_once") and CALLS[key + ":payment"] == 1:
        raise ExtractionError("scripted: every model is down")
    return PaymentOutcome(sc["payment"], "stub", 1)


async def fake_clarify_llm(model, messages):  # the clarification job also runs; keep it offline and unsent
    from schemas.disputes.clarification import ClarificationDraft
    return ClarificationDraft(subject="Your complaint", body="Hello, we are looking into it.", confidence=10), "stub", 1


import core.services.disputes.clarification as clar_mod  # noqa: E402

clar_mod.run_structured_extraction = fake_clarify_llm
intake_mod.classify_document = fake_classify
dproc.extract_dispute = fake_extract_dispute
pproc.extract_payment = fake_extract_payment


def pay(invoice, amount="900.00", ref="UTR-D3-1"):
    return ExtractedPayment(amount=D(amount), currency="INR", payment_date=date(2026, 10, 15), reference_number=ref,
                            payer_name="Zenith Softlabs Pvt Ltd", invoices=[PaidInvoice(invoice_number=invoice, amount=D(amount))],
                            confidence=PaymentFieldConfidences(amount=95, currency=95, payment_date=95, payer_name=95, reference_number=95, invoices=95))


def make_doc(path, key, body):
    d = Document()
    d.add_heading("Letter to Genworx Technologies", 1)
    d.add_paragraph(f"Reference {key}")
    d.add_paragraph(body)
    d.save(path)


async def main():
    import testsupport
    worker = await testsupport.start_worker()
    tmp = Path(tempfile.mkdtemp())

    async with SessionLocal() as db:
        zen = Customer(customer_name="Zenith Softlabs Pvt Ltd", contact="A", email="ap@zenith.com", gstin="33AABCZ1234F1Z5",
                       deducts_tds=False, billing_currency="INR", country="India")
        mer = Customer(customer_name="Meridian Tech Solutions Pvt Ltd", contact="B", email="ap@meridian.com", gstin="29AABCM5678K1Z3",
                       deducts_tds=False, billing_currency="INR", country="India")
        old = Customer(customer_name="Dormant Traders Pvt Ltd", contact="C", email="ap@dormant.com", gstin="27AABCD1111A1Z1",
                       deducts_tds=False, billing_currency="INR", country="India", is_active=False)
        db.add_all([zen, mer, old])
        await db.flush()
        for cust, nums in ((zen, ["ZEN-1", "ZEN-2", "ZEN-3", "ZEN-4", "ZEN-5", "ZEN-6", "ZEN-7", "ZEN-8"]), (mer, ["MER-1"]), (old, ["OLD-1"])):
            for n in nums:
                db.add(Invoice(customer_id=cust.id, invoice_number=n, invoice_date=date(2026, 10, 1), due_date=date(2026, 10, 31),
                               invoice_currency="INR", amount_inr=D("1000"), collectible_total=D("900")))
        await db.commit()
        ZEN, MER, OLD = zen.id, mer.id, old.id

    SCN.update({
        "SCN-S1": dict(cls=cls("dispute"), dispute=disp(invoices=["ZEN-1"], name="Zenith Softlabs Pvt Ltd", category="Quality", related=False)),
        "SCN-S2": dict(cls=cls("dispute"), dispute=disp(invoices=["ZEN-2"], email="ap@zenith.com", category="Pricing", cat_conf=60)),
        "SCN-S3": dict(cls=cls("dispute"), dispute=disp(invoices=["ZEN-3"], email="ap@zenith.com", category=None)),
        "SCN-S4": dict(cls=cls("dispute"), dispute=disp(name="Nobody Traders Pvt Ltd", category="Delivery")),
        "SCN-S5": dict(cls=cls("dispute"), dispute=disp(invoices=["ZEN-4"], email="ap@meridian.com", category="Delivery")),
        "SCN-S6": dict(cls=cls("dispute"), dispute=disp(invoices=["ZEN-6", "ZEN-7"], email="ap@zenith.com", category="Pricing", related=True)),
        "SCN-S7": dict(cls=cls("dispute"), dispute=disp(invoices=["ZEN-999"], email="ap@zenith.com", category="WrongDuplicateInvoice")),
        "SCN-S8": dict(cls=cls("dispute"), dispute=disp(invoices=["OLD-1"], gstin="27AABCD1111A1Z1", category="TaxDiscrepancy")),
        "SCN-S9": dict(cls=cls("dispute"), dispute=disp(invoices=["ZEN-1"], email="ap@zenith.com", category="Delivery")),
        "SCN-S10": dict(cls=cls("payment", also_dispute=True), dispute=disp(invoices=["ZEN-5"], email="ap@zenith.com", category="Quality", related=True),
                        payment=pay("ZEN-5", ref="UTR-D3-10")),
        "SCN-S11": dict(cls=cls("dispute", also_payment=True), dispute=disp(invoices=["ZEN-8"], email="ap@zenith.com", category="Pricing", related=True),
                        payment=pay("ZEN-8", ref="UTR-D3-11")),
        "SCN-S12": dict(cls=cls("dispute", confidence=55), dispute=disp(invoices=["MER-1"], email="ap@meridian.com", category="Quality")),
        "SCN-S13": dict(cls=cls("payment", confidence=50), dispute=disp(invoices=["ZEN-8"], email="ap@zenith.com", category="Quality"),
                        payment=pay("ZEN-8", ref="UTR-D3-13")),
        "SCN-S14": dict(cls=cls("payment", also_dispute=True), dispute=disp(invoices=["MER-1"], email="ap@meridian.com", category="Delivery"),
                        payment=ExtractedPayment(amount=D("500"), currency="INR", reference_number="UTR-D3-14", payer_name="Meridian Tech Solutions Pvt Ltd",
                                                 invoices=[PaidInvoice(invoice_number="MER-1", amount=D("500"))]), fail_payment_once=True),
        "SCN-S15": dict(cls=cls("dispute"), dispute=disp(invoices=["ZEN-3"], email="ap@zenith.com", category="Other"), fail_dispute_once=True),
        "SCN-S16": dict(cls=cls("dispute"), dispute=disp(invoices=[], email="ap@zenith.com", category="Other", summary="Our statement does not match")),
    })
    files = {}
    for key in SCN:
        files[key] = tmp / f"{key}.docx"
        make_doc(files[key], key, f"We write about invoices. Marker {key}. Please look into it.")

    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t", timeout=900) as c:
        testsupport.install(c)
        H = {"Authorization": f"Bearer {(await c.post('/auth/login', json={'email': s.seed_admin_email, 'password': s.seed_admin_password})).json()['access_token']}"}
        assert (await c.post("/admin/users", headers=H, json={"name": "Fa A", "email": "fa.a@example.com", "password": "pass12345", "role": "FA"})).status_code == 201
        await asyncio.sleep(0.05)
        assert (await c.post("/admin/users", headers=H, json={"name": "Fa B", "email": "fa.b@example.com", "password": "pass12345", "role": "FA"})).status_code == 201
        async with SessionLocal() as db:
            uid = {r[0]: str(r[1]) for r in (await db.execute(text("select email, id from users"))).all()}
        FA_ID, FB_ID = uid["fa.a@example.com"], uid["fa.b@example.com"]
        ids = {}

        async def up(key):
            r = await c.post("/documents/upload", headers=H, files={"file": (f"{key}.docx", files[key].read_bytes())})
            assert r.status_code == 202, (key, r.text)
            ids[key] = r.json()["source_document_id"]
            return (await c.get(f"/documents/uploads/{ids[key]}", headers=H)).json()

        async def disputes_of(key):
            async with SessionLocal() as db:
                return (await db.execute(text(
                    "select d.id, d.status::text, d.category::text, d.customer_id, d.invoice_id, d.assigned_to, d.payment_related, "
                    "d.classification_confidence, d.summary, d.from_email, d.classified_by, d.received_at, d.created_at, d.created_by, "
                    "d.first_response_due, d.resolution_due, i.invoice_number "
                    "from disputes d left join invoices i on i.id = d.invoice_id where d.source_document_id=:s order by i.invoice_number nulls first"),
                    {"s": ids[key]})).all()

        async def messages_of(dispute_id):
            async with SessionLocal() as db:
                return (await db.execute(text("select kind::text, body, created_by from dispute_messages where dispute_id=:d order by created_at, id"), {"d": dispute_id})).all()

        async def detail(key):
            return (await c.get(f"/documents/uploads/{ids[key]}", headers=H)).json()

        # ---- S1: a clean complaint about a known invoice -----------------------------------------
        u = await up("SCN-S1")
        ds = await disputes_of("SCN-S1")
        check("S1: status 'Dispute created', kind dispute, disposition dispute_created",
              u["status"] == "Dispute created" and u["document_kind"] == "dispute" and u["disposition"] == "dispute_created", f"{u['status']} {u['document_kind']} {u['disposition']} {u.get('error')}")
        check("S1: exactly one dispute, on invoice ZEN-1, customer Zenith, Quality, Assigned", len(ds) == 1 and ds[0][16] == "ZEN-1" and ds[0][3] == ZEN and ds[0][2] == "Quality" and ds[0][1] == "Assigned", str(ds))
        check("S1: the status carries the dispute id", u["dispute_id"] == str(ds[0][0]) and u["dispute_ids"] == [str(ds[0][0])])
        check("S1: given to the least-loaded FA (A), not classified by a person, category confidence kept", ds[0][5] is not None and str(ds[0][5]) == FA_ID and ds[0][10] is None and ds[0][7] == 92, str(ds[0][5:11]))
        check("S1: received at the upload time, deadlines set", ds[0][14] is not None and ds[0][15] is not None and ds[0][11] is not None)
        check("S1: raised by the system for the uploader (created_by is the uploader)", str(ds[0][13]) == uid[s.seed_admin_email.strip().lower()])
        check("S1: not payment related", ds[0][6] is False)
        msgs = [m for m in await messages_of(ds[0][0]) if m[0] not in ("draft", "sent")]  # the clarification job adds its own
        check("S1: the complaint text is kept as the first message", [m[0] for m in msgs] == ["inbound_email"] and "Marker SCN-S1" in msgs[0][1])
        async with SessionLocal() as db:
            flagged = (await db.execute(text("select is_disputed from invoices where invoice_number='ZEN-1'"))).scalar()
        check("S1: the invoice ZEN-1 is flagged disputed", flagged is True)
        steps = [x["step_name"] for x in (await detail("SCN-S1"))["steps"]]
        check("S1: steps parse, classify, dispute_extract, dispute_identify, dispute_route", steps == ["parse", "classify", "dispute_extract", "dispute_identify", "dispute_route"], str(steps))
        d1 = await detail("SCN-S1")
        check("S1: detail shows the extracted complaint", d1["extraction"] and d1["extraction"]["category"] == "Quality" and d1["extraction"]["invoice_numbers"] == ["ZEN-1"])
        n = (await c.get("/notifications", headers={"Authorization": f"Bearer {(await c.post('/auth/login', json={'email': 'fa.a@example.com', 'password': 'pass12345'})).json()['access_token']}"})).json()["items"]
        check("S1: the FA is notified of the assignment", any(x["type"] == "Assignment" and x["entity_id"] == str(ds[0][0]) for x in n))
        check("S1: the dispute detail names the source document", (await c.get(f"/disputes/{ds[0][0]}", headers=H)).json()["source_document_id"] == ids["SCN-S1"])

        # ---- S2/S3: the reader is unsure of the category -> the classify queue -------------------------
        u = await up("SCN-S2")
        ds = await disputes_of("SCN-S2")
        check("S2: a 60% category waits in the classify queue (Raised, suggestion kept)", ds[0][1] == "Raised" and ds[0][2] == "Pricing" and ds[0][5] is None and ds[0][7] == 60, str(ds))
        check("S2: status 'Dispute needs classification', review kind dispute_classification",
              u["status"] == "Dispute needs classification" and u["disposition"] == "dispute_queued" and u["review_kind"] == "dispute_classification" and u["review_id"] == str(ds[0][0]), f"{u['status']} {u['review_kind']}")
        msgs = await messages_of(ds[0][0])
        check("S2: reader's notes say it is only 60% sure", any(m[0] == "internal_note" and "60% sure" in m[1] for m in msgs), str([m[1][:80] for m in msgs]))
        q = (await c.get("/disputes/classify-queue", headers=H)).json()
        check("S2: it is in /disputes/classify-queue", any(x["dispute_id"] == str(ds[0][0]) for x in q))
        u = await up("SCN-S3")
        ds = await disputes_of("SCN-S3")
        check("S3: no category at all -> Raised with none, noted", ds[0][1] == "Raised" and ds[0][2] is None and any("could not tell the reason" in m[1] for m in await messages_of(ds[0][0])))

        # ---- S4: unknown customer; S5: conflicting evidence ------------------------------------------
        u = await up("SCN-S4")
        ds = await disputes_of("SCN-S4")
        check("S4: an unknown sender -> a dispute with no customer or invoice, in the queue", len(ds) == 1 and ds[0][3] is None and ds[0][4] is None and ds[0][1] == "Raised", str(ds))
        check("S4: the queue entry says the customer could not be identified", any("customer could not be identified" in m[1] for m in await messages_of(ds[0][0])))
        u = await up("SCN-S5")
        ds = await disputes_of("SCN-S5")
        check("S5: the invoice belongs to Zenith but the sender is Meridian -> no customer guessed, queued", ds[0][3] is None and ds[0][4] is None and ds[0][1] == "Raised", str(ds))
        check("S5: the conflict is written for the classifier", any("conflict" in m[1].lower() for m in await messages_of(ds[0][0])))
        async with SessionLocal() as db:
            flagged = (await db.execute(text("select is_disputed from invoices where invoice_number='ZEN-4'"))).scalar()
        check("S5: the named invoice is NOT flagged (nothing was attached)", flagged is False)

        # ---- S6: two invoices -> two disputes -------------------------------------------------------------
        u = await up("SCN-S6")
        ds = await disputes_of("SCN-S6")
        check("S6: a complaint naming two invoices raises one dispute for each", [d[16] for d in ds] == ["ZEN-6", "ZEN-7"] and all(d[1] == "Assigned" for d in ds), str([(d[16], d[1]) for d in ds]))
        check("S6: both carry payment_related", all(d[6] is True for d in ds))
        check("S6: the status lists both ids", sorted(u["dispute_ids"]) == sorted(str(d[0]) for d in ds) and len(u["dispute_ids"]) == 2)
        rel = (await c.get(f"/disputes/{ds[0][0]}", headers=H)).json()["possibly_related"]
        check("S6: disputes on DIFFERENT invoices are not 'possibly related' (decision 62 is invoice-level)", rel == [])
        async with SessionLocal() as db:
            lines = [r[0] for r in (await db.execute(text("select message from notifications where type='Assignment' and entity_id=:d"), {"d": ds[1][0]})).all()]
        check("S6: so the assignment notice has no related line", lines and all("Note:" not in (m or "") for m in lines), str(lines))

        # ---- S7: an invoice number we do not have, S8: a deactivated customer ----------------------------
        u = await up("SCN-S7")
        ds = await disputes_of("SCN-S7")
        check("S7: an unknown invoice number -> a customer-level dispute for the identified customer", len(ds) == 1 and ds[0][3] == ZEN and ds[0][4] is None and ds[0][1] == "Assigned", str(ds))
        check("S7: the unknown number is noted", any("ZEN-999" in m[1] for m in await messages_of(ds[0][0])))
        u = await up("SCN-S8")
        ds = await disputes_of("SCN-S8")
        check("S8: a deactivated customer is still heard", len(ds) == 1 and ds[0][3] == OLD and ds[0][16] == "OLD-1" and ds[0][1] == "Assigned", str(ds))

        # ---- S9: a second complaint on an invoice that already has an open dispute --------------------
        u = await up("SCN-S9")
        ds9 = await disputes_of("SCN-S9")
        ds1 = await disputes_of("SCN-S1")
        check("S9: the same invoice again (a new upload) is its OWN dispute (decision 62)", len(ds9) == 1 and ds9[0][0] != ds1[0][0] and ds9[0][16] == "ZEN-1" and ds9[0][1] == "Assigned")
        rel = (await c.get(f"/disputes/{ds9[0][0]}", headers=H)).json()["possibly_related"]
        check("S9: it shows the first one as possibly related", [r["dispute_id"] for r in rel] == [str(ds1[0][0])])
        check("S9: the first is untouched (still Assigned, no extra message)", (await disputes_of("SCN-S1"))[0][1] == "Assigned" and len([m for m in await messages_of(ds1[0][0]) if m[0] not in ("draft", "sent")]) == 1)

        # ---- S10/S11: a document with a complaint AND a payment: dispute first ------------------------------
        for key in ("SCN-S10", "SCN-S11"):
            u = await up(key)
            ds = await disputes_of(key)
            async with SessionLocal() as db:
                p = (await db.execute(text("select id, state::text, created_at, customer_id from payments where source_document_id=:d"), {"d": ids[key]})).first()
            check(f"{key[4:]}: one dispute and one payment come out of the one document", len(ds) == 1 and p is not None, f"{ds} {p}")
            check(f"{key[4:]}: the dispute was raised BEFORE the payment was written", ds[0][12] < p[2], f"{ds[0][12]} {p[2]}")
            check(f"{key[4:]}: the dispute is Assigned and payment related; the invoice is flagged", ds[0][1] == "Assigned" and ds[0][6] is True)
            check(f"{key[4:]}: the upload reports the payment and lists the dispute", u["document_kind"] == "payment" and u["payment_id"] == str(p[0]) and u["dispute_ids"] == [str(ds[0][0])], f"{u['document_kind']} {u['payment_id']} {u['dispute_ids']}")
            async with SessionLocal() as db:
                why = (await db.execute(text("select review_reasons, held_amount, held_for_dispute_id from payments where id=:i"), {"i": p[0]})).first()
            check(f"{key[4:]}: the payment, meeting a dispute that holds money, is Held for it (D6)", p[1] == "Held" and why[1] == D("900.00") and str(why[2]) == str(ds[0][0]), f"{p[1]} {u['status']} {why}")
            dd = await detail(key)
            st = [x["step_name"] for x in dd["steps"]]
            check(f"{key[4:]}: steps run the dispute steps first, then the payment steps",
                  st[:5] == ["parse", "classify", "dispute_extract", "dispute_identify", "dispute_route"] and sorted(st[5:]) == ["extract", "identify", "match", "route", "validate"], str(st))
            check(f"{key[4:]}: the payment is the main extraction and the complaint is kept beside it",
                  dd["extraction"]["reference_number"] == f"UTR-D3-{key[5:]}" and dd["dispute_extraction"]["category"] in ("Quality", "Pricing"), str(dd["dispute_extraction"] and dd["dispute_extraction"]["category"]))

        # ---- S12/S13: the reader is unsure of the type -> the uploader decides ---------------------------------
        u = await up("SCN-S12")
        check("S12: a 55% verdict waits for a type decision, no dispute yet", u["review_kind"] == "document_type" and not await disputes_of("SCN-S12"), f"{u['review_kind']}")
        rv = (await c.get("/documents/type-reviews", headers=H)).json()
        rid = next(r["review_id"] for r in rv if r["source_document_id"] == ids["SCN-S12"])
        check("S12: the queue reports the suggestion as dispute", next(r for r in rv if r["review_id"] == rid)["suggested_type"] == "dispute")
        r = await c.post(f"/documents/type-reviews/{rid}/choose", headers=H, json={"doc_type": "dispute"})
        check("S12: choosing 'dispute' is accepted", r.status_code == 200 and r.json()["processing"] is True, r.text[:100])
        await testsupport.wait_idle(ids["SCN-S12"])
        ds = await disputes_of("SCN-S12")
        dd = await detail("SCN-S12")
        check("S12: it was then read as a complaint: one dispute, Assigned; classify step says human", len(ds) == 1 and ds[0][1] == "Assigned" and dd["status"] == "Dispute created" and any(x["step_name"] == "classify" and x["output_summary"].get("source") == "human" for x in dd["steps"]), f"{dd['status']} {[d[1] for d in ds]}")
        u = await up("SCN-S13")
        rv = (await c.get("/documents/type-reviews", headers=H)).json()
        rid = next(r["review_id"] for r in rv if r["source_document_id"] == ids["SCN-S13"])
        r = await c.post(f"/documents/type-reviews/{rid}/choose", headers=H, json={"doc_type": "dispute_and_payment"})
        check("S13: choosing 'dispute_and_payment' is accepted", r.status_code == 200, r.text[:100])
        await testsupport.wait_idle(ids["SCN-S13"])
        ds = await disputes_of("SCN-S13")
        dd = await detail("SCN-S13")
        async with SessionLocal() as db:
            p = (await db.execute(text("select state::text, created_at from payments where source_document_id=:d"), {"d": ids["SCN-S13"]})).first()
        check("S13: a human 'dispute and payment' choice raises the dispute, then the payment", len(ds) == 1 and p is not None and ds[0][12] < p[1] and dd["document_kind"] == "payment" and dd["dispute_ids"] == [str(ds[0][0])], f"{[d[1] for d in ds]} {p}")
        r = await c.post(f"/documents/type-reviews/{rid}/choose", headers=H, json={"doc_type": "nonsense"})
        check("a type that does not exist -> 422", r.status_code == 422)

        # ---- S14: a retry after the complaint was raised but the payment step failed ------------------------------
        u = await up("SCN-S14")
        ds = await disputes_of("SCN-S14")
        async with SessionLocal() as db:
            pays = (await db.execute(text("select state::text from payments where source_document_id=:d"), {"d": ids["SCN-S14"]})).all()
            jobs = (await db.execute(text("select state::text, attempt_count from jobs where entity_id=:d order by created_at"), {"d": ids["SCN-S14"]})).all()
        check("S14: the payment extraction failed once and was retried (2 attempts, then Succeeded)", jobs and jobs[-1][0] == "Succeeded" and jobs[-1][1] == 2, str(jobs))
        check("S14: after the retry there is exactly ONE dispute and ONE payment (nothing written twice)", len(ds) == 1 and len(pays) == 1, f"{len(ds)} {len(pays)}")
        check("S14: the retry did not run the classifier or the complaint reader again", CALLS.get("SCN-S14:classify") == 1 and CALLS.get("SCN-S14:dispute") == 1, str({k: v for k, v in CALLS.items() if "S14" in k}))

        # ---- S15: the complaint reader failed once -> retried, one dispute ------------------------------------------
        u = await up("SCN-S15")
        ds = await disputes_of("SCN-S15")
        async with SessionLocal() as db:
            jobs = (await db.execute(text("select state::text, attempt_count from jobs where entity_id=:d order by created_at"), {"d": ids["SCN-S15"]})).all()
        check("S15: a failed complaint reading is retried and then raises exactly one dispute", jobs[-1] == ("Succeeded", 2) and len(ds) == 1 and u["status"] == "Dispute created", f"{jobs} {len(ds)} {u['status']}")

        # ---- S16: no invoice quoted, a known sender -> a customer-level dispute -------------------------------------------
        u = await up("SCN-S16")
        ds = await disputes_of("SCN-S16")
        check("S16: no invoice named -> a customer-level dispute (E1), categorised and assigned", len(ds) == 1 and ds[0][3] == ZEN and ds[0][4] is None and ds[0][1] == "Assigned" and ds[0][2] == "Other", str(ds))

        # ---- the upload list and permissions --------------------------------------------------------------------------------
        lst = (await c.get("/documents/uploads", headers=H)).json()
        check("the uploads list shows complaints with their kind", sum(1 for x in lst if x["document_kind"] == "dispute") >= 10, str([x["document_kind"] for x in lst][:5]))
        check("uploading needs a token", (await c.post("/documents/upload", files={"file": ("x.docx", b"x")})).status_code == 401)

        # ---- housekeeping ---------------------------------------------------------------------------------------------------------
        async with SessionLocal() as db:
            found = await verify_money_identities(db)
            dup = (await db.execute(text("select count(*) from (select source_document_id from disputes where source_document_id is not null group by source_document_id, invoice_id having count(*) > 1) t"))).scalar()
        check("the money checker holds", found == [], str(found))
        check("no document raised the same invoice's dispute twice", dup == 0)

    await worker.stop()
    await engine.dispose()
    print(f"RESULT {sum(results)}/{len(results)} passed")


asyncio.run(main())
