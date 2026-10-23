import asyncio
import re
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, "src")
import httpx
from docx import Document
from sqlalchemy import text

import core.services.documents.intake as intake
from api.rest.app import app
from config.settings import get_settings
from core.services.documents.classification import ClassificationOutcome
from data.clients.db import SessionLocal, engine
from schemas.documents.classification import DocClassification, DocumentClass

s = get_settings()
assert "ar_sb" in s.database_url, "refusing to run against a non-test database"
OURS = "33AAHCG2210A1Z7"
results = []


def check(name, cond, extra=""):
    results.append(bool(cond))
    print(f"RESULT {'PASS' if cond else 'FAIL'} {name} {extra}"[:320])


async def classify(markdown):
    m = re.search(r"FORCE:(\w+):(\d+)", markdown)
    verdict = DocClassification(doc_class=DocumentClass(m.group(1)), confidence=int(m.group(2)), reason="forced for the test")
    return ClassificationOutcome(verdict, "fake-model", 1)


intake.classify_document = classify


def invoice(path, number, marker, due="Due Date: 04/11/2026"):
    d = Document()
    d.add_heading("TAX INVOICE", 1)
    d.add_paragraph(f"Genworx Technologies Pvt Ltd\nGSTIN: {OURS}")
    d.add_paragraph(f"Invoice No: {number}\nInvoice Date: 05/10/2026\n{due}")
    d.add_paragraph("Bill To:\nZenith Softlabs Pvt Ltd\nGSTIN: 33AABCZ1234F1Z5")
    t = d.add_table(rows=2, cols=3)
    t.style = "Table Grid"
    for i, h in enumerate(("Description", "Qty", "Amount (INR)")):
        t.rows[0].cells[i].text = h
    for i, v in enumerate(("Consulting", "1", "50000.00")):
        t.rows[1].cells[i].text = v
    d.add_paragraph("Total INR 50000.00")
    d.add_paragraph(marker)
    d.save(path)


async def main():
    import testsupport
    worker = await testsupport.start_worker()
    tmp = Path(tempfile.mkdtemp())
    invoice(tmp / "wrong_other.docx", "INV-SB-1", "FORCE:other:95")
    invoice(tmp / "human_rejected.docx", "INV-SB-2", "FORCE:unclear:80")
    invoice(tmp / "nodue.docx", "INV-SB-3", "FORCE:invoice:95", due="")
    invoice(tmp / "twice.docx", "INV-SB-4", "FORCE:other:90")

    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t", timeout=600) as c:
        testsupport.install(c)
        login = lambda e, p: c.post("/auth/login", json={"email": e, "password": p})
        ADM = {"Authorization": f"Bearer {(await login(s.seed_admin_email, s.seed_admin_password)).json()['access_token']}"}
        await c.post("/admin/users", headers=ADM, json={"name": "Fa Two", "email": "fa.two@example.com", "password": "pass12345", "role": "FA"})
        FA = {"Authorization": f"Bearer {(await login('fa.two@example.com', 'pass12345')).json()['access_token']}"}
        cust = lambda **kw: {"contact": "Contact", "country": "India", "billing_currency": "INR", **kw}
        assert (await c.post("/customers", headers=ADM, json=cust(customer_name="Zenith Softlabs Pvt Ltd", email="ap@zenith.com", gstin="33AABCZ1234F1Z5", deducts_tds=True))).status_code == 201

        async def q(sql, **p):
            async with SessionLocal() as db:
                return (await db.execute(text(sql), p)).all()

        async def wait(sid, n=1):
            for _ in range(240):
                js = await q("select state from jobs where entity_id=:i order by created_at", i=sid)
                if len(js) >= n and js[-1][0] in ("Succeeded", "Failed"):
                    return
                await asyncio.sleep(0.5)

        async def up(name, H=ADM):
            r = await c.post("/documents/upload", headers=H, files={"file": (name, (tmp / name).read_bytes())})
            assert r.status_code == 202, r.text
            sid = r.json()["source_document_id"]
            await wait(sid)
            return sid

        async def status(sid, H=ADM):
            r = await c.get(f"/documents/uploads/{sid}", headers=H)
            return r.json() if r.status_code == 200 else (await c.get(f"/documents/uploads/{sid}", headers=H)).json()

        async def rejected(H=ADM, **params):
            return (await c.get("/invoices/rejected", params=params, headers=H)).json()

        # ---- a confident wrong 'other'
        sid = await up("wrong_other.docx", H=ADM)
        item = next(x for x in await rejected(kind="other") if x["source_document_id"] == sid)
        check("rejected 'neither' entry offers can_send_to_review", item["can_send_to_review"] is True and item["file_kept"] is True and item["kind"] == "other", str(item["can_send_to_review"]))
        rid = item["rejected_id"]

        r = await c.post(f"/invoices/rejected/{rid}/send-to-review", headers=FA)
        check("another user -> 403", r.status_code == 403, r.text[:100])
        check("...and nothing changed", (await q("select resolution::text from rejected_invoices where id=:i", i=rid))[0][0] == "None")
        r = await c.post("/invoices/rejected/00000000-0000-0000-0000-000000000009/send-to-review", headers=ADM)
        check("unknown id -> 404", r.status_code == 404)
        check("no token -> 401", (await c.post(f"/invoices/rejected/{rid}/send-to-review")).status_code == 401)

        rs = await asyncio.gather(
            c.post(f"/invoices/rejected/{rid}/send-to-review", headers=ADM),
            c.post(f"/invoices/rejected/{rid}/send-to-review", headers=ADM),
        )
        check("two simultaneous send-backs: one 200, one 409", sorted(x.status_code for x in rs) == [200, 409], str([x.status_code for x in rs]))
        ok = next(x for x in rs if x.status_code == 200).json()
        check("response carries the new review id", ok["rejected_id"] == rid and ok["review_id"], str(ok))
        check("exactly one new review opened", (await q("select count(*) from pending_invoice_changes where source_document_id=:i and proposed_snapshot->>'kind'='doc_type'", i=sid))[0][0] == 1)
        check("rejection marked resolved (Reuploaded) and hidden by default", (await q("select resolution::text from rejected_invoices where id=:i", i=rid))[0][0] == "Reuploaded" and rid not in [x["rejected_id"] for x in await rejected(kind="other")] and rid in [x["rejected_id"] for x in await rejected(kind="other", include_resolved="true")])
        shown = next(x for x in await rejected(kind="other", include_resolved="true") if x["rejected_id"] == rid)
        check("resolved entry no longer offers the button", shown["can_send_to_review"] is False)
        check("send-back logged on the document", (await q("select count(*) from state_log where entity_id=:i and to_state='SentToTypeReview'", i=sid))[0][0] == 1)
        r = await c.post(f"/invoices/rejected/{rid}/send-to-review", headers=ADM)
        check("sending it back again -> 409", r.status_code == 409, r.text[:100])

        queue = (await c.get("/documents/type-reviews", headers=ADM)).json()
        review = next((x for x in queue if x["review_id"] == ok["review_id"]), None)
        check("the type-review queue shows it, with the original reason and the sent-back note", review and any("sent this back" in x for x in review["reasons"]) and any("neither an invoice" in x for x in review["reasons"]), str(review and review["reasons"]))
        check("status says Needs review again", (await status(sid))["status"] == "Needs review")
        r = await c.post(f"/documents/type-reviews/{ok['review_id']}/choose", headers=ADM, json={"doc_type": "invoice", "reason": "It is an invoice"})
        await wait(sid, 2)
        st = await status(sid)
        check("choosing invoice runs the normal invoice flow -> Invoice created", r.status_code == 200 and st["status"] == "Invoice created" and st["invoice_number"] == "INV-SB-1", f"{st['status']} {st.get('disposition_reasons')}")

        # ---- a human rejection can be taken back too
        sid2 = await up("human_rejected.docx", H=FA)
        rev = (await status(sid2, H=FA))["review_id"]
        await c.post(f"/documents/type-reviews/{rev}/reject", headers=FA, json={"reason": "Rejected by mistake"})
        item2 = next(x for x in await rejected(H=FA, kind="other") if x["source_document_id"] == sid2)
        check("a rejection made in the type review also offers the button", item2["can_send_to_review"] is True)
        r = await c.post(f"/invoices/rejected/{item2['rejected_id']}/send-to-review", headers=FA)
        check("...and it can be sent back by its uploader", r.status_code == 200, r.text[:100])
        r2 = await c.post(f"/documents/type-reviews/{r.json()['review_id']}/reject", headers=FA, json={"reason": "No, really reject"})
        check("it can then be rejected again (new rejected entry, file kept)", r2.status_code == 200 and len([x for x in await rejected(H=FA, kind="other") if x["source_document_id"] == sid2]) == 1 and (await q("select file_path is not null from source_documents where id=:i", i=sid2))[0][0])

        # ---- only 'neither' rejections qualify
        sid3 = await up("nodue.docx")
        inv_item = next(x for x in await rejected() if x["source_document_id"] == sid3)
        check("a validation rejection does not offer the button", inv_item["can_send_to_review"] is False and inv_item["kind"] == "invoice")
        r = await c.post(f"/invoices/rejected/{inv_item['rejected_id']}/send-to-review", headers=ADM)
        check("...and sending it is refused (422)", r.status_code == 422, r.text[:120])

        # ---- file gone
        sid4 = await up("twice.docx")
        item4 = next(x for x in await rejected(kind="other") if x["source_document_id"] == sid4)
        async with SessionLocal() as db:
            await db.execute(text("update source_documents set file_path=null where id=:i"), {"i": sid4})
            await db.commit()
        r = await c.post(f"/invoices/rejected/{item4['rejected_id']}/send-to-review", headers=ADM)
        check("file no longer on disk -> 422, nothing opened", r.status_code == 422 and (await q("select count(*) from pending_invoice_changes where source_document_id=:i", i=sid4))[0][0] == 0, r.text[:100])
        check("file-less entry does not offer the button", next(x for x in await rejected(kind="other") if x["source_document_id"] == sid4)["can_send_to_review"] is False)
    await worker.stop(grace_seconds=5)
    await engine.dispose()
    print(f"RESULT {sum(results)}/{len(results)} passed")


asyncio.run(main())
